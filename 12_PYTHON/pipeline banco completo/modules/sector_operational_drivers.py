"""
Drivers operacionais por setor para enriquecer os workbooks de valuation.

Esses dados normalmente nao aparecem na CVM de forma padronizada. A fonte
esperada costuma ser release de resultados, apresentacao de RI, formulario de
referencia, relatorio da administracao ou planilhas internas do analista.
"""

from __future__ import annotations

import csv
import re
import unicodedata
from pathlib import Path
from typing import Any

from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from modules.excel_theme import THEME


def _slug(text: str) -> str:
    text = unicodedata.normalize("NFKD", str(text or "").strip().lower())
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")


KPI_CATALOG: dict[str, list[dict[str, str]]] = {
    "bancos": [
        {"driver": "Carteira de credito", "unidade": "R$ MM", "fonte": "CVM/DFP + release", "uso": "Base para NIM, provisao e crescimento do FCFE."},
        {"driver": "NIM / margem financeira", "unidade": "%", "fonte": "CVM + release", "uso": "Principal driver da margem financeira bruta."},
        {"driver": "Inadimplencia >90 dias", "unidade": "%", "fonte": "release/apresentacao RI", "uso": "Ancora PDD/PCLD e risco de credito."},
        {"driver": "Indice de eficiencia", "unidade": "%", "fonte": "release/apresentacao RI", "uso": "Valida despesas operacionais vs margem financeira."},
        {"driver": "Basileia / capital principal", "unidade": "%", "fonte": "release/Bacen", "uso": "Limita crescimento, payout e reinvestimento regulatorio."},
    ],
    "varejo": [
        {"driver": "Numero de lojas", "unidade": "lojas", "fonte": "release/apresentacao RI", "uso": "Separa crescimento por expansao fisica vs produtividade."},
        {"driver": "Area de vendas", "unidade": "m2", "fonte": "release/apresentacao RI", "uso": "Permite receita por m2 e ganho de produtividade."},
        {"driver": "Vendas mesmas lojas (SSS)", "unidade": "%", "fonte": "release/apresentacao RI", "uso": "Driver organico de receita sem abertura de lojas."},
        {"driver": "Receita por loja", "unidade": "R$ MM/loja", "fonte": "calculado", "uso": "Check de coerencia da projecao de receita."},
        {"driver": "Giro de estoque", "unidade": "x", "fonte": "CVM + release", "uso": "Ancora capital de giro e risco de markdown."},
        {"driver": "Margem bruta", "unidade": "%", "fonte": "CVM/release", "uso": "Ponte entre receita, mix, promocao e EBITDA."},
    ],
    "energia": [
        {"driver": "Energia vendida / distribuida", "unidade": "GWh", "fonte": "release/apresentacao RI", "uso": "Driver de volume e receita."},
        {"driver": "RAP / receita regulatoria", "unidade": "R$ MM", "fonte": "ANEEL/release", "uso": "Ancora receita de transmissoras."},
        {"driver": "Capex regulatorio", "unidade": "R$ MM", "fonte": "release/guidance", "uso": "Base de investimento e reinvestimento."},
        {"driver": "Disponibilidade / perdas", "unidade": "%", "fonte": "release/regulador", "uso": "Afeta eficiencia e penalidades."},
    ],
    "saneamento": [
        {"driver": "Economias / ligacoes", "unidade": "mil", "fonte": "release/apresentacao RI", "uso": "Driver de volume e capilaridade."},
        {"driver": "Volume faturado", "unidade": "m3", "fonte": "release/apresentacao RI", "uso": "Ancora receita operacional."},
        {"driver": "Indice de perdas", "unidade": "%", "fonte": "release/regulador", "uso": "Check operacional de margem e capex."},
        {"driver": "Tarifa media", "unidade": "R$/m3", "fonte": "calculado/release", "uso": "Separa volume de preco/regulacao."},
    ],
    "industrial": [
        {"driver": "Volume vendido / produzido", "unidade": "unid./ton", "fonte": "release/apresentacao RI", "uso": "Separa crescimento por volume e preco."},
        {"driver": "Preco medio / mix", "unidade": "R$/unid.", "fonte": "release/calculado", "uso": "Explica margem e receita."},
        {"driver": "Utilizacao de capacidade", "unidade": "%", "fonte": "release", "uso": "Valida alavancagem operacional."},
        {"driver": "Backlog / carteira", "unidade": "R$ MM", "fonte": "release", "uso": "Sustenta previsibilidade de receita."},
    ],
    "petroleo": [
        {"driver": "Producao media", "unidade": "boe/d", "fonte": "release/ANP", "uso": "Principal driver de receita."},
        {"driver": "Brent medio", "unidade": "US$/bbl", "fonte": "mercado", "uso": "Driver de preco e sensibilidade."},
        {"driver": "Lifting cost", "unidade": "US$/boe", "fonte": "release", "uso": "Ancora margem operacional."},
        {"driver": "Reservas 2P / vida util", "unidade": "boe/anos", "fonte": "relatorio reservas", "uso": "Sustenta valor terminal e risco de reposicao."},
    ],
    "locadoras": [
        {"driver": "Frota media", "unidade": "veiculos", "fonte": "release/apresentacao RI", "uso": "Base de receita e capex."},
        {"driver": "Taxa de utilizacao", "unidade": "%", "fonte": "release", "uso": "Check de eficiencia operacional."},
        {"driver": "Ticket diario / yield", "unidade": "R$/dia", "fonte": "release/calculado", "uso": "Driver de receita por ativo."},
        {"driver": "Spread compra-venda", "unidade": "%", "fonte": "release", "uso": "Afeta margem de seminovos e depreciação economica."},
    ],
    "saude": [
        {"driver": "Ticket medio", "unidade": "R$", "fonte": "release/apresentacao RI", "uso": "Separa preco de volume."},
        {"driver": "Volume de exames/procedimentos", "unidade": "mil", "fonte": "release", "uso": "Driver de receita."},
        {"driver": "Ocupacao / sinistralidade", "unidade": "%", "fonte": "release/regulador", "uso": "Valida margem e risco setorial."},
        {"driver": "Unidades / leitos", "unidade": "qtd.", "fonte": "release/formulario ref.", "uso": "Mede capacidade instalada."},
    ],
}


SECTOR_ALIASES = {
    "consumo_varejo": "varejo",
    "consumo": "varejo",
    "retail": "varejo",
    "banco": "bancos",
    "bank": "bancos",
    "oleo_gas": "petroleo",
    "oil_gas": "petroleo",
}


def normalizar_setor_driver(setor: str | None, tipo_empresa: str | None = None) -> str:
    if tipo_empresa == "bank":
        return "bancos"
    key = _slug(setor or "geral")
    return SECTOR_ALIASES.get(key, key)


def catalogo_por_setor(setor: str | None, tipo_empresa: str | None = None) -> list[dict[str, str]]:
    key = normalizar_setor_driver(setor, tipo_empresa)
    return KPI_CATALOG.get(key, [])


def carregar_drivers_operacionais(
    ticker: str,
    setor: str | None,
    tipo_empresa: str | None,
    root_dir: str | Path,
) -> dict[str, Any]:
    """Carrega CSV opcional data/operational_drivers/TICKER.csv."""
    root_dir = Path(root_dir)
    key = normalizar_setor_driver(setor, tipo_empresa)
    catalogo = catalogo_por_setor(setor, tipo_empresa)
    path = root_dir / "data" / "operational_drivers" / f"{ticker.upper()}.csv"
    valores: dict[str, dict[int, Any]] = {}
    fontes: dict[str, str] = {}

    if path.exists():
        with path.open("r", encoding="utf-8-sig", newline="") as fh:
            reader = csv.DictReader(fh)
            for row in reader:
                driver = row.get("driver") or row.get("indicador") or row.get("metric")
                ano_raw = row.get("ano") or row.get("year")
                valor = row.get("valor") or row.get("value")
                if not driver or not ano_raw:
                    continue
                try:
                    ano = int(float(str(ano_raw).strip()))
                except Exception:
                    continue
                slug = _slug(driver)
                try:
                    valor_final: Any = float(str(valor).replace(",", "."))
                except Exception:
                    valor_final = valor
                valores.setdefault(slug, {})[ano] = valor_final
                fonte = row.get("fonte") or row.get("source") or ""
                if fonte:
                    fontes[slug] = fonte

    return {
        "setor_driver": key,
        "catalogo": catalogo,
        "valores": valores,
        "fontes": fontes,
        "source_path": str(path) if path.exists() else "",
    }


def adicionar_aba_drivers_setoriais(wb, dados: dict, nome_empresa: str, ticker: str) -> None:
    """Adiciona uma aba de contexto operacional setorial ao workbook."""
    if "Drivers Setoriais" in wb.sheetnames:
        del wb["Drivers Setoriais"]

    meta = (dados or {}).get("metodologia", {})
    tipo = meta.get("tipo_empresa")
    setor = meta.get("setor")
    drivers = (dados or {}).get("drivers_operacionais") or {}
    catalogo = drivers.get("catalogo") or catalogo_por_setor(setor, tipo)
    anos = list((meta.get("premissas_efetivas") or {}).get("anos_historicos") or [])
    anos += list((meta.get("premissas_efetivas") or {}).get("anos_projecao") or [])[:5]

    ws = wb.create_sheet("Drivers Setoriais")
    max_col = 5 + max(len(anos), 1)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    title = ws.cell(row=1, column=1, value=f"Drivers Setoriais - {nome_empresa} ({ticker})")
    title.font = THEME.header_font(13)
    title.fill = THEME.fill(THEME.HEADER)
    title.alignment = Alignment(horizontal="center")

    ws.cell(row=3, column=1, value="Objetivo")
    ws.cell(
        row=3,
        column=2,
        value="Separar dados contabeis da CVM de KPIs operacionais de RI que melhoram a leitura setorial.",
    )
    ws.cell(row=4, column=1, value="Fonte esperada")
    ws.cell(row=4, column=2, value="Releases, apresentacoes, formulario de referencia, guidance ou CSV manual em data/operational_drivers/TICKER.csv.")

    headers = ["Driver operacional", "Unidade", "Fonte esperada", "Uso no valuation", "Status"] + anos
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=6, column=col, value=header)
        cell.font = THEME.header_font(9)
        cell.fill = THEME.fill(THEME.HEADER)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    valores = drivers.get("valores") or {}
    fontes = drivers.get("fontes") or {}
    for row_idx, item in enumerate(catalogo, start=7):
        driver = item.get("driver", "")
        slug = _slug(driver)
        serie = valores.get(slug, {})
        status = "OK" if serie else "Pendente RI/manual"
        fonte = fontes.get(slug) or item.get("fonte", "")
        base = [driver, item.get("unidade", ""), fonte, item.get("uso", ""), status]
        for col, value in enumerate(base, start=1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.font = THEME.font(color=THEME.TEXT, bold=(col == 1), size=9)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=row_idx, column=5).fill = THEME.fill(THEME.ACCENT if serie else THEME.PROJECTION_BG)
        for offset, ano in enumerate(anos, start=6):
            value = serie.get(int(ano))
            cell = ws.cell(row=row_idx, column=offset, value=value if value is not None else "-")
            cell.font = THEME.font(color=THEME.HIST if value is not None else THEME.TEXT, size=9)
            cell.alignment = Alignment(horizontal="right")
            if isinstance(value, (int, float)):
                cell.number_format = THEME.NUM_2

    if not catalogo:
        ws.cell(row=7, column=1, value="Sem catalogo setorial especifico cadastrado.")
        ws.cell(row=7, column=2, value="Adicionar drivers em modules/sector_operational_drivers.py.")

    widths = {1: 28, 2: 12, 3: 28, 4: 44, 5: 18}
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    for col in range(6, max_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 12
    ws.freeze_panes = "F7"
