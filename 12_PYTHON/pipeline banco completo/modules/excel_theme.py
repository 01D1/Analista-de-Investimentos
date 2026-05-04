"""
Tema visual padrao dos workbooks de valuation.

Este modulo centraliza a identidade visual. A metodologia de valuation continua
separada nos motores de projecao/valuation e nas premissas por setor/empresa.
"""

from openpyxl.styles import Font, PatternFill, Border, Side


class ValuationTheme:
    # Cores principais
    HIST = "0000FF"
    PROJ = "008000"
    TEXT = "000000"
    WHITE = "FFFFFF"
    HEADER = "1F4E79"
    SUBHEADER = "D6E4F0"
    ACCENT = "E2EFDA"
    PROJECTION_BG = "FFF2CC"
    NAV_BG = "F2F2F2"
    NEGATIVE = "FF0000"
    WARNING = "ED7D31"
    GOLD = "FFC000"
    DARK_GREEN = "375623"
    BORDER = "D9D9D9"
    FONT = "Calibri"

    # Formatos numéricos
    NUM = '#,##0;(#,##0);"-"'
    NUM_1 = '#,##0.0;(#,##0.0);"-"'
    NUM_2 = '#,##0.00;(#,##0.00);"-"'
    PCT_2 = '0.00%;(0.00%);"-"'
    PCT_1 = '0.0%;(0.0%);"-"'
    BRL = 'R$ #,##0.00;(R$ #,##0.00);"-"'
    MULT = '0.0x;(0.0x);"-"'

    @classmethod
    def font(cls, color: str | None = None, bold: bool = False, size: int = 9) -> Font:
        return Font(color=color or cls.TEXT, bold=bold, size=size, name=cls.FONT)

    @classmethod
    def header_font(cls, size: int = 10) -> Font:
        return Font(color=cls.WHITE, bold=True, size=size, name=cls.FONT)

    @staticmethod
    def fill(color: str) -> PatternFill:
        return PatternFill(start_color=color, end_color=color, fill_type="solid")

    @classmethod
    def border(cls) -> Border:
        side = Side(style="thin", color=cls.BORDER)
        return Border(bottom=side, top=side, left=side, right=side)


THEME = ValuationTheme

