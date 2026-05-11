"""
post_excel_quality_gate.py

Valida a planilha final depois da escrita no Excel. A ideia aqui nao e
substituir o modelo financeiro, mas impedir que um valuation seja marcado como
concluido quando a pasta de trabalho saiu com abas vazias, referencias quebradas,
anos incoerentes ou celulas criticas sem valor.
"""

from __future__ import annotations

import json
import logging
import re
import unicodedata
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

logger = logging.getLogger("pipeline.post_excel_quality")


FORMULA_ERROR_TOKENS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NULL!", "#NUM!")


def _norm_text(value: str) -> str:
    value = unicodedata.normalize("NFKD", str(value or "").lower())
    return "".join(ch for ch in value if not unicodedata.combining(ch))


@dataclass
class ExcelFinding:
    severity: str
    sheet: str
    cell: str
    issue: str
    impact: str
    suggestion: str


def _cell_label(cell) -> str:
    return getattr(cell, "coordinate", "-")


def _is_empty(value: Any) -> bool:
    return value is None or str(value).strip() in {"", "-", "None"}


def _sheet_non_empty_cells(ws, max_rows: int = 220, max_cols: int = 80) -> int:
    count = 0
    for row in ws.iter_rows(
        min_row=1,
        max_row=min(ws.max_row or 1, max_rows),
        min_col=1,
        max_col=min(ws.max_column or 1, max_cols),
    ):
        for cell in row:
            if not _is_empty(cell.value):
                count += 1
    return count


def _find_year_headers(ws) -> list[int]:
    candidates: list[list[int]] = []
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, 12)):
        row_years = []
        for cell in row:
            value = cell.value
            if isinstance(value, int) and 1990 <= value <= 2100:
                row_years.append(value)
            elif isinstance(value, str) and re.fullmatch(r"20\d{2}", value.strip()):
                row_years.append(int(value.strip()))
        if len(row_years) >= 2:
            candidates.append(row_years)
    if not candidates:
        return []
    return max(candidates, key=len)


def _scan_formula_errors(ws) -> list[ExcelFinding]:
    findings: list[ExcelFinding] = []
    for row in ws.iter_rows():
        for cell in row:
            value = cell.value
            if isinstance(value, str) and any(token in value.upper() for token in FORMULA_ERROR_TOKENS):
                findings.append(ExcelFinding(
                    severity="critica",
                    sheet=ws.title,
                    cell=_cell_label(cell),
                    issue=f"Formula ou valor contem erro: {value}",
                    impact="Pode quebrar o valuation final ou esconder uma premissa sem vinculo.",
                    suggestion="Corrigir a referencia/fonte antes de confiar no preco justo.",
                ))
    return findings


def _scan_blank_sheets(wb) -> list[ExcelFinding]:
    findings: list[ExcelFinding] = []
    for ws in wb.worksheets:
        filled = _sheet_non_empty_cells(ws)
        if filled < 5:
            findings.append(ExcelFinding(
                severity="alta",
                sheet=ws.title,
                cell="A1",
                issue="Aba praticamente vazia.",
                impact="O analista pode estar vendo uma aba esperada sem dados, formulas ou premissas.",
                suggestion="Revisar escritor Excel e mapeamento da aba.",
            ))
    return findings


def _scan_year_consistency(wb, anos_hist: list[int], anos_proj: list[int]) -> list[ExcelFinding]:
    findings: list[ExcelFinding] = []
    expected = set(int(a) for a in (anos_hist or []) + (anos_proj or []))
    if not expected:
        return findings
    for ws in wb.worksheets:
        years = _find_year_headers(ws)
        if not years:
            continue
        duplicates = sorted({year for year in years if years.count(year) > 1})
        if duplicates:
            findings.append(ExcelFinding(
                severity="media",
                sheet=ws.title,
                cell="cabecalho",
                issue=f"Anos duplicados no cabecalho: {duplicates}",
                impact="Pode haver sobrescrita ou desalinhamento entre historico e projecao.",
                suggestion="Regerar cabecalho da aba e conferir mapa de colunas.",
            ))
        missing_hist = [a for a in anos_hist if a not in years]
        if missing_hist and any(token in ws.title.lower() for token in ["dre", "balan", "balanço", "dcf", "fluxo"]):
            findings.append(ExcelFinding(
                severity="media",
                sheet=ws.title,
                cell="cabecalho",
                issue=f"Historico ausente no cabecalho: {missing_hist}",
                impact="A planilha pode estar omitindo anos historicos que alimentam premissas.",
                suggestion="Conferir anos_historicos usados pelo escritor Excel.",
            ))
    return findings


def _scan_required_labels(wb, tipo_empresa: str) -> list[ExcelFinding]:
    required = ["valor do equity", "preco justo", "upside"]
    if tipo_empresa == "bank":
        required += ["lucro liquido", "fcfe", "ke"]
    else:
        required += ["receita", "ebit", "fcff", "wacc"]

    all_text = {}
    for ws in wb.worksheets:
        labels = []
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, 220), max_col=min(ws.max_column or 1, 20)):
            for cell in row:
                if isinstance(cell.value, str):
                    labels.append(_norm_text(cell.value))
        all_text[ws.title] = " | ".join(labels)

    joined = "\n".join(all_text.values())
    findings: list[ExcelFinding] = []
    for label in required:
        if label not in joined:
            findings.append(ExcelFinding(
                severity="alta",
                sheet="workbook",
                cell="-",
                issue=f"Linha/rotulo critico nao encontrado: {label}",
                impact="Pode indicar que o layout gerado nao esta refletindo a metodologia do valuation.",
                suggestion="Adicionar ou mapear a linha critica no escritor Excel.",
            ))
    return findings


def _is_zeroish(value: Any) -> bool:
    if isinstance(value, bool) or _is_empty(value):
        return False
    try:
        return abs(float(value)) < 1e-12
    except (TypeError, ValueError):
        return False


def _scan_dashboard_integrity(wb, tipo_empresa: str) -> list[ExcelFinding]:
    findings: list[ExcelFinding] = []
    if "Dashboard" not in wb.sheetnames:
        return findings

    ws = wb["Dashboard"]
    max_row = min(ws.max_row or 1, 90)
    max_col = min(ws.max_column or 1, 30)

    for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
        for cell in row:
            value = cell.value
            if not isinstance(value, str) or not value.startswith("="):
                continue
            normalized = value.replace(" ", "").upper()
            if "!X85" in normalized or "!X88" in normalized:
                findings.append(ExcelFinding(
                    severity="alta",
                    sheet="Dashboard",
                    cell=_cell_label(cell),
                    issue=f"Formula de resumo aponta para linha/coluna antiga: {value}",
                    impact="O dashboard pode exibir upside, TIR ou preco teto herdado de layout antigo.",
                    suggestion="Regerar a planilha com o escritor atual ou sobrescrever o campo pelo dado calculado.",
                ))

    labels: list[tuple[str, int, int]] = []
    for row in range(1, max_row + 1):
        for col in (4, 7):
            label = ws.cell(row=row, column=col).value
            if isinstance(label, str):
                labels.append((_norm_text(label), row, col))

    def find_cell(*needles: str):
        normalized_needles = [_norm_text(needle) for needle in needles]
        for label, row, col in labels:
            if all(needle in label for needle in normalized_needles):
                return ws.cell(row=row, column=col + 1)
        return None

    checked_cells: set[str] = set()

    def check_value(*needles: str, severity: str, impact: str, zero_is_issue: bool = True) -> None:
        cell = find_cell(*needles)
        if cell is None:
            return
        if cell.coordinate in checked_cells:
            return
        checked_cells.add(cell.coordinate)
        value = cell.value
        if _is_empty(value) or (zero_is_issue and _is_zeroish(value)):
            findings.append(ExcelFinding(
                severity=severity,
                sheet="Dashboard",
                cell=_cell_label(cell),
                issue=f"Campo critico do dashboard sem valor valido: {' / '.join(needles)}",
                impact=impact,
                suggestion="Conferir o mapeamento entre mercado, valuation e Dashboard.",
            ))

    check_value("Cotacao", "ON", severity="alta", impact="Cotacao zerada distorce upside e leitura da recomendacao.")
    check_value("Valor do Equity", severity="alta", impact="Equity value vazio ou zerado exige revisao do valuation.")
    check_value("Preco Justo", "ON", severity="alta", impact="Preco justo ON vazio ou zerado exige revisao do valuation.")
    check_value("Upside", "ON", severity="media", impact="Upside ON vazio ou zerado pode esconder falha de vinculo.")
    check_value("TIR", "ON", severity="media", impact="TIR ON vazia ou zerada reduz a confiabilidade do resumo.")
    check_value("Preco Teto", "ON", severity="alta", impact="Preco teto ON vazio ou zerado exige revisao do valuation.")

    if tipo_empresa == "bank":
        check_value("Acoes", "ON", severity="media", impact="Base acionaria ON zerada distorce preco justo por classe.")
        check_value("Cotacao", "PN", severity="alta", impact="Cotacao PN zerada distorce upside e leitura da recomendacao.")
        check_value("Preco Justo", "PN", severity="alta", impact="Preco justo PN vazio ou zerado exige revisao do valuation.")
        check_value("Upside", "PN", severity="media", impact="Upside PN vazio ou zerado pode esconder falha de vinculo.")
        check_value("TIR", "PN", severity="media", impact="TIR PN vazia ou zerada reduz a confiabilidade do resumo.")
        check_value("Preco Teto", "PN", severity="alta", impact="Preco teto PN vazio ou zerado exige revisao do valuation.")

    return findings


def executar_quality_gate_excel(
    *,
    ticker: str,
    excel_path: str | Path,
    tipo_empresa: str,
    anos_hist: list[int],
    anos_proj: list[int],
    output_dir: str | Path,
) -> dict[str, Any]:
    excel_path = Path(excel_path)
    out_dir = Path(output_dir) / "post_excel_quality"
    out_dir.mkdir(parents=True, exist_ok=True)
    findings: list[ExcelFinding] = []

    try:
        wb = load_workbook(excel_path, data_only=False, read_only=True)
    except Exception as exc:
        findings.append(ExcelFinding(
            severity="critica",
            sheet="workbook",
            cell="-",
            issue=f"Nao foi possivel abrir o Excel: {exc}",
            impact="Valuation final nao auditavel.",
            suggestion="Fechar o Excel se estiver aberto e regerar o arquivo.",
        ))
        wb = None

    sheet_count = 0
    if wb is not None:
        sheet_count = len(wb.worksheets)
        findings.extend(_scan_blank_sheets(wb))
        findings.extend(_scan_year_consistency(wb, anos_hist, anos_proj))
        findings.extend(_scan_required_labels(wb, tipo_empresa))
        findings.extend(_scan_dashboard_integrity(wb, tipo_empresa))
        for ws in wb.worksheets:
            findings.extend(_scan_formula_errors(ws))
        wb.close()

    severity_rank = {"baixa": 1, "media": 2, "alta": 3, "critica": 4}
    critical = [f for f in findings if f.severity == "critica"]
    high = [f for f in findings if f.severity == "alta"]
    score = max(0, 100 - len(critical) * 35 - len(high) * 18 - sum(8 for f in findings if f.severity == "media"))
    status = "bloqueado" if critical else "revisar" if high or score < 80 else "aprovado"

    result = {
        "ticker": ticker,
        "timestamp": datetime.now().isoformat(),
        "excel_path": str(excel_path),
        "sheet_count": sheet_count,
        "status": status,
        "score": score,
        "bloqueia": bool(critical),
        "findings": [asdict(f) for f in sorted(findings, key=lambda f: severity_rank.get(f.severity, 0), reverse=True)],
    }

    json_path = out_dir / f"post_excel_quality_{ticker}.json"
    md_path = out_dir / f"post_excel_quality_{ticker}.md"
    json_path.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    lines = [
        f"# Quality Gate Final - {ticker}",
        "",
        f"- Status: **{status}**",
        f"- Score: **{score}/100**",
        f"- Arquivo: `{excel_path}`",
        "",
        "## Achados",
    ]
    if findings:
        for f in result["findings"]:
            lines.append(f"- **{f['severity']}** | {f['sheet']}!{f['cell']} | {f['issue']} Impacto: {f['impact']} Correção: {f['suggestion']}")
    else:
        lines.append("- Nenhum problema critico encontrado.")
    md_path.write_text("\n".join(lines), encoding="utf-8")
    result["json_path"] = str(json_path)
    result["markdown_path"] = str(md_path)
    logger.info("[post_excel_quality] %s: %s score=%s findings=%d", ticker, status, score, len(findings))
    return result
