"""
qualitative_engine.py

Motor qualitativo para organizar documentos corporativos, extrair evidencias,
classificar eventos e gerar uma ficha qualitativa por empresa.

Principio central: quando nao ha documento ou evidencia, o modulo registra
"sem evidencia" em vez de inferir informacao ausente.
"""

from __future__ import annotations

import hashlib
import csv
import io
import json
import logging
import re
import shutil
import zipfile
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Optional
from urllib.request import Request, urlopen

try:
    import requests
except Exception:  # pragma: no cover
    requests = None

logger = logging.getLogger("pipeline.qualitative")

CVM_IPE_DADOS_URL = "https://dados.cvm.gov.br/dados/CIA_ABERTA/DOC/IPE/DADOS/ipe_cia_aberta_{ano}.zip"
USER_AGENT = "PipelineValuationQualitative/1.0"


def _http_get(url: str, timeout: int = 30) -> tuple[bytes, dict[str, str]]:
    if requests is not None:
        resp = requests.get(url, timeout=timeout, headers={"User-Agent": USER_AGENT})
        resp.raise_for_status()
        return resp.content, dict(resp.headers)
    req = Request(url, headers={"User-Agent": USER_AGENT})
    with urlopen(req, timeout=timeout) as resp:  # nosec: URL controlada por fontes publicas/CVM
        headers = {k: v for k, v in resp.headers.items()}
        return resp.read(), headers


QUAL_DIRS = {
    "base": Path("data/qualitative"),
    "raw": Path("data/qualitative/raw"),
    "processed": Path("data/qualitative/processed"),
    "summaries": Path("data/qualitative/summaries"),
    "events": Path("data/qualitative/events"),
    "reports": Path("data/qualitative/reports"),
}

SUPPORTED_EXTS = {".txt", ".md", ".html", ".htm", ".pdf", ".docx", ".json", ".csv"}

DOC_TYPE_KEYWORDS = {
    "release_resultados": ["release", "resultado", "earnings", "itr", "trimestre"],
    "apresentacao": ["apresentacao", "presentation", "ri", "institucional"],
    "dfp_itr": ["itr", "dfp", "demonstracoes financeiras", "informacoes trimestrais"],
    "notas_explicativas": ["notas explicativas", "nota explicativa"],
    "formulario_referencia": ["formulario de referencia", "fre", "referencia"],
    "fato_relevante": ["fato relevante", "material fact"],
    "comunicado_mercado": ["comunicado ao mercado", "comunicado"],
    "ata": ["ata", "assembleia", "conselho"],
    "relatorio_administracao": ["relatorio da administracao", "administracao"],
    "guidance": ["guidance", "projecoes", "estimativa", "meta"],
    "teleconferencia": ["call", "teleconferencia", "transcricao", "transcript"],
}

EVENT_PATTERNS = {
    "mudanca_gestao": ["renuncia", "eleicao", "diretor", "ceo", "cfo", "presidente", "conselho"],
    "estrategia": ["estrategia", "plano", "reposicionamento", "transformacao", "turnaround"],
    "aquisicao_desinvestimento": ["aquisicao", "fusao", "incorporacao", "venda de ativo", "desinvestimento"],
    "guidance": ["guidance", "meta", "projecao", "estimativa"],
    "risco_juridico": ["processo", "contingencia", "litigio", "arbitragem", "acao civil"],
    "risco_financeiro": ["endividamento", "liquidez", "covenant", "rating", "alavancagem"],
    "risco_operacional": ["interrupcao", "acidente", "parada", "falha operacional", "fornecedor"],
    "risco_regulatorio": ["regulatorio", "aneel", "anatel", "banco central", "cade", "licenca"],
    "risco_reputacional": ["investigacao", "fraude", "corrupcao", "reputacional", "sancao"],
    "produto_operacao": ["novo produto", "lancamento", "capacidade", "expansao", "loja", "planta"],
}

POSITIVE_WORDS = [
    "crescimento", "melhora", "expansao", "ganho", "recorde", "reduziu custos",
    "aumento de margem", "desalavancagem", "forte demanda", "aprovado",
]
NEGATIVE_WORDS = [
    "queda", "reducao", "perda", "deterioracao", "pressao", "risco", "atraso",
    "provisao", "contingencia", "inadimplencia", "investigacao", "renuncia",
]
CRITICAL_WORDS = ["fraude", "recuperacao judicial", "covenant", "intervencao", "corrupcao", "suspensao"]
HIGH_WORDS = ["fato relevante", "guidance", "aquisicao", "desinvestimento", "renuncia", "contingencia"]

SCORE_DIMENSIONS = [
    "governanca",
    "transparencia",
    "previsibilidade",
    "qualidade_gestao",
    "execucao_estrategica",
    "risco_regulatorio",
    "risco_financeiro",
    "risco_operacional",
    "risco_reputacional",
    "alinhamento_minoritarios",
    "consistencia_discurso_numeros",
]

DIMENSION_RULES = {
    "governanca": ["conselho", "comite", "governanca", "assembleia", "independente"],
    "transparencia": ["transparencia", "divulgacao", "comunicado", "release", "apresentacao"],
    "previsibilidade": ["guidance", "previsibilidade", "recorrente", "estavel", "contrato"],
    "qualidade_gestao": ["gestao", "diretoria", "ceo", "cfo", "execucao"],
    "execucao_estrategica": ["estrategia", "plano", "execucao", "meta", "expansao"],
    "risco_regulatorio": ["regulatorio", "licenca", "aneel", "anatel", "banco central"],
    "risco_financeiro": ["endividamento", "liquidez", "covenant", "rating", "alavancagem"],
    "risco_operacional": ["operacional", "producao", "parada", "acidente", "fornecedor"],
    "risco_reputacional": ["reputacao", "fraude", "corrupcao", "investigacao", "sancao"],
    "alinhamento_minoritarios": ["dividendos", "tag along", "minoritarios", "payout", "recompra"],
    "consistencia_discurso_numeros": ["guidance", "aderencia", "entrega", "meta", "resultado"],
}


@dataclass
class QualitativeDocument:
    ticker: str
    source_path: str
    preserved_path: str
    document_id: str
    document_type: str
    title: str
    date: Optional[str] = None
    source_url: Optional[str] = None
    text_path: Optional[str] = None
    char_count: int = 0
    collected_at: str = field(default_factory=lambda: datetime.now().isoformat())


@dataclass
class QualitativeEvent:
    ticker: str
    event_type: str
    title: str
    description: str
    relevance: str
    thesis_effect: str
    source_document: str
    evidence: str
    date: Optional[str] = None


@dataclass
class DimensionScore:
    score: Optional[float]
    scale: str
    confidence: str
    evidence_count: int
    rationale: str
    sources: list[str]


def ensure_qualitative_dirs(base_dir: Path | str | None = None) -> dict[str, Path]:
    root = Path(base_dir) if base_dir else Path(".")
    paths = {name: root / rel for name, rel in QUAL_DIRS.items()}
    for path in paths.values():
        path.mkdir(parents=True, exist_ok=True)
    return paths


def _safe_name(value: str) -> str:
    value = re.sub(r"[^\w.-]+", "_", str(value or "").strip(), flags=re.UNICODE)
    return value.strip("_")[:160] or "documento"


def _hash_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()[:16]


def _clean_text(text: str) -> str:
    text = re.sub(r"\s+", " ", text or "")
    return text.strip()


def _is_usable_text(text: str) -> bool:
    text = text or ""
    if len(text.strip()) < 40:
        return False
    sample = text[:5000]
    printable = sum(1 for ch in sample if ch.isprintable() or ch.isspace())
    letters = sum(1 for ch in sample if ch.isalpha())
    return printable / max(len(sample), 1) >= 0.90 and letters / max(len(sample), 1) >= 0.10


def _extract_text_pdf(path: Path) -> str:
    try:
        import pdfplumber
        parts = []
        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages[:80]:
                text = page.extract_text() or ""
                if text:
                    parts.append(text)
        return "\n".join(parts)
    except Exception:
        try:
            import fitz
            doc = fitz.open(str(path))
            return "\n".join(page.get_text("text") for page in doc[:80])
        except Exception as exc:
            logger.debug("[qualitative] Falha ao extrair PDF %s: %s", path, exc)
            return ""


def _extract_text_docx(path: Path) -> str:
    try:
        import docx
        doc = docx.Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs if p.text)
    except Exception:
        try:
            with zipfile.ZipFile(path) as zf:
                xml = zf.read("word/document.xml").decode("utf-8", errors="ignore")
            xml = re.sub(r"<[^>]+>", " ", xml)
            return _clean_text(xml)
        except Exception as exc:
            logger.debug("[qualitative] Falha ao extrair DOCX %s: %s", path, exc)
            return ""


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix in {".txt", ".md", ".csv", ".json"}:
            return path.read_text(encoding="utf-8", errors="ignore")
        if suffix in {".html", ".htm"}:
            html = path.read_text(encoding="utf-8", errors="ignore")
            html = re.sub(r"<script[^>]*>.*?</script>", " ", html, flags=re.I | re.S)
            html = re.sub(r"<style[^>]*>.*?</style>", " ", html, flags=re.I | re.S)
            text = _clean_text(re.sub(r"<[^>]+>", " ", html))
            return text if _is_usable_text(text) else ""
        if suffix == ".pdf":
            return _extract_text_pdf(path)
        if suffix == ".docx":
            return _extract_text_docx(path)
    except Exception as exc:
        logger.debug("[qualitative] Falha ao extrair texto %s: %s", path, exc)
    return ""


def classify_document(path: Path, text: str = "") -> str:
    haystack = f"{path.name} {text[:3000]}".lower()
    haystack = _remove_accents(haystack)
    scores = {}
    for doc_type, words in DOC_TYPE_KEYWORDS.items():
        scores[doc_type] = sum(1 for word in words if _remove_accents(word) in haystack)
    best, score = max(scores.items(), key=lambda kv: kv[1])
    return best if score > 0 else "outro"


def _remove_accents(value: str) -> str:
    import unicodedata

    value = unicodedata.normalize("NFKD", value or "")
    return "".join(ch for ch in value if not unicodedata.combining(ch))


def _sentences(text: str, max_sentences: int = 800) -> list[str]:
    raw = re.split(r"(?<=[.!?])\s+|\n+", text or "")
    out = []
    for sentence in raw:
        sentence = _clean_text(sentence)
        if re.match(r"^(fonte|link|link_download|link_doc|codigo|cd_cvm|cnpj|versao)\s*:", sentence, flags=re.I):
            continue
        if 40 <= len(sentence) <= 600:
            out.append(sentence)
        if len(out) >= max_sentences:
            break
    return out


def discover_local_documents(ticker: str, paths: dict[str, Path]) -> list[Path]:
    candidates: list[Path] = []
    ticker = ticker.upper()
    roots = [paths["raw"] / ticker, paths["raw"]]
    for root in roots:
        if not root.exists():
            continue
        for path in root.rglob("*"):
            if path.is_file() and path.suffix.lower() in SUPPORTED_EXTS:
                if root.name.upper() == ticker or ticker.lower() in path.name.lower():
                    candidates.append(path)
    return sorted(set(candidates), key=lambda p: (p.stat().st_mtime, str(p)), reverse=True)


def _parse_cvm_csv_from_zip(content: bytes) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with zipfile.ZipFile(io.BytesIO(content)) as zf:
        for name in zf.namelist():
            if not name.lower().endswith(".csv"):
                continue
            raw = zf.read(name)
            text = raw.decode("latin1", errors="ignore")
            reader = csv.DictReader(io.StringIO(text), delimiter=";")
            rows.extend(dict(row) for row in reader)
    return rows


def _download_cvm_ipe_year(ano: int, paths: dict[str, Path]) -> list[dict[str, Any]]:
    cache_dir = paths["processed"] / "_cvm_ipe_cache"
    cache_dir.mkdir(parents=True, exist_ok=True)
    cache_path = cache_dir / f"ipe_cia_aberta_{ano}.zip"
    try:
        if cache_path.exists() and cache_path.stat().st_size > 0:
            content = cache_path.read_bytes()
        else:
            url = CVM_IPE_DADOS_URL.format(ano=ano)
            content, _headers = _http_get(url, timeout=30)
            cache_path.write_bytes(content)
        return _parse_cvm_csv_from_zip(content)
    except Exception as exc:
        logger.info("[qualitative] IPE CVM %s indisponivel: %s", ano, exc)
        return []


def _row_get(row: dict[str, Any], *names: str) -> str:
    for name in names:
        if name in row and row.get(name) not in (None, ""):
            return str(row.get(name)).strip()
    low = {str(k).lower(): k for k in row.keys()}
    for name in names:
        key = low.get(name.lower())
        if key and row.get(key) not in (None, ""):
            return str(row.get(key)).strip()
    return ""


def _cvm_row_date(row: dict[str, Any]) -> str:
    return _row_get(row, "DT_REFER", "DT_RECEB", "DT_ENTREGA", "DT_DOC", "DATA_REFERENCIA")


def _cvm_row_link(row: dict[str, Any]) -> str:
    for key, value in row.items():
        key_norm = str(key).lower()
        value_str = str(value or "").strip()
        if value_str.startswith("http") and any(token in key_norm for token in ["link", "url", "doc"]):
            return value_str
    return ""


def _cvm_doc_priority(row: dict[str, Any]) -> int:
    text = _remove_accents(" ".join(str(v or "") for v in row.values()).lower())
    priorities = [
        ("fato relevante", 100),
        ("comunicado ao mercado", 90),
        ("dados economico-financeiros", 85),
        ("assembleia", 80),
        ("reuniao da administracao", 75),
        ("politica de dividendos", 70),
        ("politica de gerenciamento de riscos", 70),
        ("relatorio de sustentabilidade", 65),
        ("codigo de conduta", 60),
        ("calendario de eventos corporativos", 55),
    ]
    for token, score in priorities:
        if token in text:
            return score
    return 20


def _cvm_row_to_text(row: dict[str, Any], link: str = "") -> str:
    lines = [
        "Fonte: CVM - Documentos Periodicos e Eventuais (IPE)",
        f"Link: {link or 'sem link direto no indice'}",
        "",
    ]
    for key in sorted(row.keys()):
        value = str(row.get(key, "") or "").strip()
        if value:
            lines.append(f"{key}: {value}")
    return "\n".join(lines)


def _save_downloaded_content(dest_base: Path, content: bytes, content_type: str, fallback_text: str) -> Optional[Path]:
    content_type = content_type.lower()
    if content.startswith(b"%PDF") or "pdf" in content_type:
        suffix = ".pdf"
    elif content.startswith(b"PK\x03\x04"):
        dest_base.with_suffix(".bin").write_bytes(content)
        return None
    else:
        decoded = content[:5000].decode("utf-8", errors="ignore")
        if "<html" in decoded.lower() or _is_usable_text(decoded):
            suffix = ".html"
        else:
            dest_base.with_suffix(".bin").write_bytes(content)
            return None
    dest = dest_base.with_suffix(suffix)
    if len(content) >= 100:
        dest.write_bytes(content)
    else:
        dest = dest_base.with_suffix(".txt")
        dest.write_text(fallback_text, encoding="utf-8")
    return dest


def auto_collect_cvm_documents(
    ticker: str,
    codigo_cvm: str | None,
    paths: dict[str, Path],
    max_docs: int = 20,
    years_back: int = 5,
) -> list[Path]:
    """Busca automaticamente documentos IPE/CVM e salva evidencias em raw/<TICKER>/auto_cvm."""
    if not codigo_cvm:
        return []

    ticker = ticker.upper()
    out_dir = paths["raw"] / ticker / "auto_cvm"
    out_dir.mkdir(parents=True, exist_ok=True)
    current_year = datetime.now().year
    rows: list[dict[str, Any]] = []
    for ano in range(current_year, current_year - years_back - 1, -1):
        for row in _download_cvm_ipe_year(ano, paths):
            if str(_row_get(row, "CD_CVM", "Codigo_CVM", "COD_CVM")).strip() == str(codigo_cvm).strip():
                rows.append(row)

    rows = sorted(
        rows,
        key=lambda r: (_cvm_doc_priority(r), _cvm_row_date(r), _row_get(r, "VERSAO", "CD_DOC")),
        reverse=True,
    )[:max_docs]

    collected: list[Path] = []
    for idx, row in enumerate(rows, start=1):
        category = _row_get(row, "CATEG_DOC", "CATEGORIA", "DS_CATEGORIA", "Tipo")
        date = _safe_name(_cvm_row_date(row) or f"doc_{idx:02d}")
        title = _safe_name(_row_get(row, "ASSUNTO", "DS_ASSUNTO", "DESC_ASSUNTO", "CATEG_DOC", "ID_DOC") or category or f"documento_{idx:02d}")
        base = out_dir / f"{idx:02d}_{date}_{title}"
        link = _cvm_row_link(row)
        metadata_text = _cvm_row_to_text(row, link)
        metadata_path = base.with_suffix(".txt")
        metadata_path.write_text(metadata_text, encoding="utf-8")
        collected.append(metadata_path)

        if not link:
            continue
        try:
            content, headers = _http_get(link, timeout=30)
            downloaded = _save_downloaded_content(base.with_name(base.name + "_documento"), content, headers.get("content-type", ""), metadata_text)
            if downloaded is not None:
                collected.append(downloaded)
        except Exception as exc:
            logger.info("[qualitative] Nao foi possivel baixar documento CVM %s: %s", link[:120], exc)

    index_path = paths["processed"] / ticker / "cvm_auto_collect_index.json"
    index_path.parent.mkdir(parents=True, exist_ok=True)
    index_path.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    logger.info("[qualitative] CVM/IPE auto: %s docs/metadata salvos para %s", len(collected), ticker)
    return collected


def preserve_document(ticker: str, source: Path, paths: dict[str, Path]) -> Path:
    doc_hash = _hash_file(source)
    dest_dir = paths["raw"] / ticker.upper()
    dest_dir.mkdir(parents=True, exist_ok=True)
    dest = dest_dir / f"{doc_hash}_{_safe_name(source.name)}"
    if source.resolve() != dest.resolve() and not dest.exists():
        shutil.copy2(source, dest)
    return dest


def process_documents(ticker: str, paths: dict[str, Path], local_docs: Iterable[Path]) -> list[QualitativeDocument]:
    processed: list[QualitativeDocument] = []
    seen_hashes: set[str] = set()
    out_dir = paths["processed"] / ticker.upper()
    out_dir.mkdir(parents=True, exist_ok=True)
    for source in local_docs:
        preserved = preserve_document(ticker, source, paths)
        doc_id = _hash_file(preserved)
        if doc_id in seen_hashes:
            continue
        seen_hashes.add(doc_id)
        text = extract_text(preserved)
        if not _is_usable_text(text):
            continue
        doc_type = classify_document(preserved, text)
        text_path = out_dir / f"{doc_id}.txt"
        meta_path = out_dir / f"{doc_id}.json"
        if text:
            text_path.write_text(text, encoding="utf-8")
        doc = QualitativeDocument(
            ticker=ticker.upper(),
            source_path=str(source),
            preserved_path=str(preserved),
            document_id=doc_id,
            document_type=doc_type,
            title=source.stem,
            text_path=str(text_path) if text else None,
            char_count=len(text),
        )
        meta_path.write_text(json.dumps(asdict(doc), ensure_ascii=False, indent=2), encoding="utf-8")
        processed.append(doc)
    return processed


def collect_cvm_index(ticker: str, codigo_cvm: str | None, paths: dict[str, Path]) -> list[dict[str, Any]]:
    """Coleta indice CVM/IPE dos ultimos anos. Se a rede falhar, nao bloqueia."""
    if not codigo_cvm:
        return []
    index_dir = paths["processed"] / ticker.upper()
    index_dir.mkdir(parents=True, exist_ok=True)
    docs = []
    current_year = datetime.now().year
    for ano in range(current_year, current_year - 6, -1):
        for row in _download_cvm_ipe_year(ano, paths):
            if str(_row_get(row, "CD_CVM", "Codigo_CVM", "COD_CVM")).strip() == str(codigo_cvm).strip():
                docs.append(row)
    docs = sorted(docs, key=lambda r: (_cvm_row_date(r), _cvm_doc_priority(r)), reverse=True)[:100]
    (index_dir / "cvm_ipe_index.json").write_text(json.dumps(docs, ensure_ascii=False, indent=2), encoding="utf-8")
    return docs


def _event_relevance(sentence: str) -> str:
    text = _remove_accents(sentence.lower())
    if any(word in text for word in CRITICAL_WORDS):
        return "critica"
    if any(word in text for word in HIGH_WORDS):
        return "alta"
    if any(word in text for word in NEGATIVE_WORDS + POSITIVE_WORDS):
        return "media"
    return "baixa"


def _thesis_effect(sentence: str) -> str:
    text = _remove_accents(sentence.lower())
    pos = sum(1 for word in POSITIVE_WORDS if word in text)
    neg = sum(1 for word in NEGATIVE_WORDS + CRITICAL_WORDS if word in text)
    if pos > neg:
        return "melhora"
    if neg > pos:
        return "piora"
    return "mantem"


def extract_events(ticker: str, docs: list[QualitativeDocument], paths: dict[str, Path]) -> list[QualitativeEvent]:
    events: list[QualitativeEvent] = []
    seen = set()
    for doc in docs:
        if not doc.text_path:
            continue
        text = Path(doc.text_path).read_text(encoding="utf-8", errors="ignore")
        for sentence in _sentences(text):
            low = _remove_accents(sentence.lower())
            for event_type, words in EVENT_PATTERNS.items():
                if not any(_remove_accents(w) in low for w in words):
                    continue
                key = hashlib.sha1(f"{event_type}|{sentence[:200]}".encode("utf-8")).hexdigest()
                if key in seen:
                    continue
                seen.add(key)
                events.append(QualitativeEvent(
                    ticker=ticker.upper(),
                    event_type=event_type,
                    title=event_type.replace("_", " ").title(),
                    description=sentence[:450],
                    relevance=_event_relevance(sentence),
                    thesis_effect=_thesis_effect(sentence),
                    source_document=doc.title,
                    evidence=sentence[:600],
                    date=doc.date,
                ))
    events = sorted(events, key=lambda e: ["critica", "alta", "media", "baixa"].index(e.relevance))
    out_dir = paths["events"] / ticker.upper()
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "events.json").write_text(
        json.dumps([asdict(e) for e in events], ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return events


def build_scorecard(ticker: str, docs: list[QualitativeDocument], events: list[QualitativeEvent],
                    valuation_context: dict[str, Any] | None = None) -> dict[str, Any]:
    evidence_by_dim: dict[str, list[tuple[str, str]]] = {dim: [] for dim in SCORE_DIMENSIONS}
    for doc in docs:
        if not doc.text_path:
            continue
        text = Path(doc.text_path).read_text(encoding="utf-8", errors="ignore")
        low = _remove_accents(text.lower())
        for dim, words in DIMENSION_RULES.items():
            for word in words:
                w = _remove_accents(word)
                idx = low.find(w)
                if idx >= 0:
                    snippet = _clean_text(text[max(0, idx - 180): idx + 260])
                    evidence_by_dim[dim].append((doc.title, snippet))
                    break

    event_by_type = {}
    for event in events:
        event_by_type.setdefault(event.event_type, []).append(event)

    scores: dict[str, DimensionScore] = {}
    for dim in SCORE_DIMENSIONS:
        evidence = evidence_by_dim.get(dim, [])
        if not evidence:
            scores[dim] = DimensionScore(
                score=None,
                scale="0-5",
                confidence="sem_evidencia",
                evidence_count=0,
                rationale="Nao ha documento suficiente para pontuar sem inferencia.",
                sources=[],
            )
            continue

        relevant_events = [
            e for e in events
            if dim.replace("risco_", "risco_") in e.event_type or any(w in _remove_accents(e.evidence.lower()) for w in DIMENSION_RULES.get(dim, []))
        ]
        base = 3.0
        positives = sum(1 for e in relevant_events if e.thesis_effect == "melhora")
        negatives = sum(1 for e in relevant_events if e.thesis_effect == "piora")
        if dim.startswith("risco_"):
            score = base - min(2.0, negatives * 0.5) + min(1.0, positives * 0.25)
        else:
            score = base + min(1.5, positives * 0.4) - min(1.5, negatives * 0.4)
        score = max(0.0, min(5.0, round(score, 1)))
        confidence = "alta" if len(evidence) >= 5 else "media" if len(evidence) >= 2 else "baixa"
        scores[dim] = DimensionScore(
            score=score,
            scale="0-5",
            confidence=confidence,
            evidence_count=len(evidence),
            rationale="Pontuacao heuristica baseada em evidencias documentais e eventos extraidos.",
            sources=sorted({src for src, _ in evidence})[:10],
        )

    scored_values = [s.score for s in scores.values() if s.score is not None]
    overall = round(sum(scored_values) / len(scored_values), 2) if scored_values else None
    alert = build_thesis_alert(events, valuation_context or {})

    return {
        "ticker": ticker.upper(),
        "generated_at": datetime.now().isoformat(),
        "overall_score": overall,
        "scale": "0-5",
        "score_policy": "Sem evidencia documental suficiente, a dimensao fica como null em vez de nota inventada.",
        "dimensions": {dim: asdict(score) for dim, score in scores.items()},
        "events_count": len(events),
        "documents_count": len(docs),
        "thesis_alert": alert,
    }


def build_thesis_alert(events: list[QualitativeEvent], valuation_context: dict[str, Any]) -> dict[str, Any]:
    critical = [e for e in events if e.relevance == "critica"]
    high_negative = [e for e in events if e.relevance == "alta" and e.thesis_effect == "piora"]
    high_positive = [e for e in events if e.relevance == "alta" and e.thesis_effect == "melhora"]
    if critical or len(high_negative) >= 2:
        level = "alto"
        message = "Eventos qualitativos podem exigir revisao imediata da tese."
    elif high_negative:
        level = "medio"
        message = "Ha evento negativo relevante para monitorar antes de manter premissas."
    elif high_positive:
        level = "baixo"
        message = "Ha catalisador positivo, mas sem indicacao automatica de mudanca de tese."
    else:
        level = "baixo"
        message = "Nenhuma mudanca qualitativa relevante detectada nos documentos analisados."
    return {
        "level": level,
        "message": message,
        "critical_events": len(critical),
        "negative_high_events": len(high_negative),
        "positive_high_events": len(high_positive),
        "valuation_reference": {
            "preco_justo_on": valuation_context.get("preco_justo_on"),
            "upside_on": valuation_context.get("upside_on"),
            "tir_on": valuation_context.get("tir_on"),
        },
    }


def _top_events(events: list[QualitativeEvent], n: int = 8) -> list[QualitativeEvent]:
    return events[:n]


def _collect_event_text(events: list[QualitativeEvent], effect: str | None = None, event_type_prefix: str | None = None) -> list[str]:
    out = []
    for event in events:
        if effect and event.thesis_effect != effect:
            continue
        if event_type_prefix and not event.event_type.startswith(event_type_prefix):
            continue
        out.append(f"{event.description} Fonte: {event.source_document}.")
    return out[:8]


def _bullets(items: list[str], empty: str = "Sem evidencia documental suficiente.") -> str:
    if not items:
        return f"- {empty}"
    return "\n".join(item if str(item).lstrip().startswith("- ") else f"- {item}" for item in items)


def append_run_log(ticker: str, result: dict[str, Any], paths: dict[str, Path]) -> None:
    log_path = paths["base"] / "qualitative_runs.log"
    payload = {
        "timestamp": datetime.now().isoformat(),
        "ticker": ticker.upper(),
        "documents_count": result.get("documents_count"),
        "events_count": result.get("events_count"),
        "overall_score": result.get("overall_score"),
        "report_path": result.get("report_path"),
        "scorecard_path": result.get("scorecard_path"),
        "cvm_index_count": result.get("cvm_index_count"),
        "ri_documents_count": result.get("ri_documents_count"),
    }
    with log_path.open("a", encoding="utf-8") as fh:
        fh.write(json.dumps(payload, ensure_ascii=False) + "\n")


def write_report(ticker: str, nome: str, docs: list[QualitativeDocument], events: list[QualitativeEvent],
                 scorecard: dict[str, Any], paths: dict[str, Path],
                 valuation_context: dict[str, Any] | None = None) -> Path:
    out_dir = paths["reports"] / ticker.upper()
    out_dir.mkdir(parents=True, exist_ok=True)
    report_path = out_dir / "RELATORIO_QUALITATIVO_EMPRESA.md"
    latest_path = paths["reports"] / f"RELATORIO_QUALITATIVO_{ticker.upper()}_latest.md"

    docs_lines = [
        f"- {doc.title} | tipo={doc.document_type} | caracteres={doc.char_count} | fonte=`{doc.preserved_path}`"
        for doc in docs[:20]
    ]
    score_lines = []
    for dim, data in scorecard.get("dimensions", {}).items():
        value = data.get("score")
        value_txt = "sem evidencia" if value is None else f"{value}/5"
        score_lines.append(f"- {dim}: {value_txt} | confianca={data.get('confidence')} | fontes={', '.join(data.get('sources', [])[:3]) or '-'}")

    positives = _collect_event_text(events, effect="melhora")
    negatives = _collect_event_text(events, effect="piora")
    risks = [e.description + f" Fonte: {e.source_document}." for e in events if e.event_type.startswith("risco")][:10]
    catalysts = [e.description + f" Fonte: {e.source_document}." for e in events if e.thesis_effect == "melhora"][:10]
    management = [e.description + f" Fonte: {e.source_document}." for e in events if e.event_type == "mudanca_gestao"][:10]
    strategy = [e.description + f" Fonte: {e.source_document}." for e in events if e.event_type == "estrategia"][:10]

    valuation_context = valuation_context or {}
    impact = []
    if risks:
        impact.append("Riscos qualitativos podem justificar desconto maior, WACC/Ke mais conservador ou reducao de crescimento terminal.")
    if catalysts:
        impact.append("Catalisadores positivos podem sustentar maior crescimento, margem ou reducao de risco se confirmados nos proximos resultados.")
    if not impact:
        impact.append("Sem evidencia suficiente para alterar premissas numericas do valuation.")

    content = f"""# RELATORIO_QUALITATIVO_EMPRESA

Empresa: **{nome} ({ticker.upper()})**  
Gerado em: {datetime.now():%Y-%m-%d %H:%M:%S}

## Resumo executivo

- Documentos analisados: **{len(docs)}**.
- Eventos extraidos: **{len(events)}**.
- Score qualitativo geral: **{scorecard.get('overall_score') if scorecard.get('overall_score') is not None else 'sem evidencia suficiente'}**.
- Alerta de tese: **{scorecard.get('thesis_alert', {}).get('level')}** - {scorecard.get('thesis_alert', {}).get('message')}

## Documentos usados

{_bullets(docs_lines, "Nenhum documento local encontrado em data/qualitative/raw.")}

## Principais fatos recentes

{_bullets([f"{e.relevance.upper()} | {e.thesis_effect}: {e.description} Fonte: {e.source_document}." for e in _top_events(events)])}

## Leitura critica dos releases

### Pontos positivos
{_bullets(positives)}

### Pontos negativos
{_bullets(negatives)}

## Mudancas internas relevantes

{_bullets(management)}

## Mudancas recentes na estrategia

{_bullets(strategy)}

## Riscos nao capturados diretamente pelos numeros

{_bullets(risks)}

## Oportunidades nao capturadas diretamente pelos numeros

{_bullets(catalysts)}

## Score qualitativo

{_bullets(score_lines)}

## Guidance e aderencia historica

- Esta primeira versao registra eventos e evidencias de guidance quando aparecem nos documentos, mas nao calcula aderencia historica automaticamente sem uma base de guidance estruturada.

## Temas recorrentes em notas explicativas

- Sem classificacao recorrente automatica quando nao houver notas explicativas processadas com texto extraido.

## Impacto provavel sobre premissas de valuation

{_bullets(impact)}

Referencia numerica disponivel: preco justo ON={valuation_context.get('preco_justo_on')}, upside ON={valuation_context.get('upside_on')}, TIR ON={valuation_context.get('tir_on')}.

## Perguntas que o analista ainda deveria investigar

- Quais eventos extraidos possuem impacto financeiro quantificavel nas premissas de receita, margem, CAPEX, capital de giro ou WACC?
- O guidance divulgado foi cumprido nos ultimos trimestres?
- Houve mudanca de diretoria, conselho ou controlador que altere risco de execucao?
- Existem contingencias juridicas ou regulatorias relevantes nas notas explicativas?
- A narrativa dos releases e consistente com os numeros normalizados no valuation?

## Conclusao qualitativa

{scorecard.get('thesis_alert', {}).get('message')}

> Regra de confiabilidade: conclusoes acima usam apenas documentos listados como fonte. Ausencia de documento nao foi convertida em opiniao.
"""
    report_path.write_text(content, encoding="utf-8")
    latest_path.write_text(content, encoding="utf-8")
    return report_path


def run_qualitative_analysis(
    ticker: str,
    nome: str | None = None,
    codigo_cvm: str | None = None,
    base_dir: Path | str | None = None,
    valuation_context: dict[str, Any] | None = None,
    collect_cvm: bool = False,
    ri_urls: Iterable[str] | str | None = None,
    collect_ri: bool = True,
    ri_max_docs: int = 25,
    ri_depth: int = 1,
) -> dict[str, Any]:
    ticker = ticker.upper()
    nome = nome or ticker
    paths = ensure_qualitative_dirs(base_dir)
    logger.info("[qualitative] Iniciando analise qualitativa de %s", ticker)

    cvm_index = collect_cvm_index(ticker, codigo_cvm, paths) if collect_cvm else []
    auto_docs = auto_collect_cvm_documents(ticker, codigo_cvm, paths) if collect_cvm else []
    ri_result = None
    if collect_ri and ri_urls:
        try:
            from modules.ri_crawler import crawl_ri_sources

            ri_result = crawl_ri_sources(
                ticker=ticker,
                source_urls=ri_urls,
                base_dir=base_dir or ".",
                max_docs=ri_max_docs,
                depth=ri_depth,
            )
        except Exception as exc:
            logger.warning("[qualitative] Falha no crawler RI de %s: %s", ticker, exc)
    local_docs = discover_local_documents(ticker, paths)
    docs = process_documents(ticker, paths, local_docs)
    events = extract_events(ticker, docs, paths)
    scorecard = build_scorecard(ticker, docs, events, valuation_context)

    out_summary = paths["summaries"] / ticker
    out_summary.mkdir(parents=True, exist_ok=True)
    scorecard_path = out_summary / "qualitative_scorecard.json"
    scorecard["cvm_index_count"] = len(cvm_index)
    scorecard["ri_crawler"] = ri_result
    scorecard["ri_documents_count"] = int((ri_result or {}).get("saved_documents") or 0)
    scorecard["documents"] = [asdict(doc) for doc in docs]
    scorecard["events"] = [asdict(event) for event in events]
    scorecard_path.write_text(json.dumps(scorecard, ensure_ascii=False, indent=2), encoding="utf-8")

    report_path = write_report(ticker, nome, docs, events, scorecard, paths, valuation_context)

    result = {
        "ticker": ticker,
        "documents_count": len(docs),
        "events_count": len(events),
        "overall_score": scorecard.get("overall_score"),
        "thesis_alert": scorecard.get("thesis_alert"),
        "scorecard_path": str(scorecard_path),
        "report_path": str(report_path),
        "cvm_index_count": len(cvm_index),
        "auto_collected_documents_count": len(auto_docs),
        "ri_documents_count": int((ri_result or {}).get("saved_documents") or 0),
        "ri_crawler": ri_result,
    }
    append_run_log(ticker, result, paths)

    logger.info(
        "[qualitative] %s concluido: docs=%d eventos=%d score=%s",
        ticker, len(docs), len(events), scorecard.get("overall_score"),
    )
    return result


if __name__ == "__main__":  # pragma: no cover
    import argparse

    parser = argparse.ArgumentParser(description="Motor qualitativo de empresas")
    parser.add_argument("--ticker", required=True)
    parser.add_argument("--nome", default=None)
    parser.add_argument("--codigo-cvm", default=None)
    parser.add_argument("--collect-cvm", action="store_true")
    parser.add_argument("--ri-url", action="append", default=None)
    parser.add_argument("--sem-ri-crawler", action="store_true")
    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(name)s: %(message)s")
    result = run_qualitative_analysis(
        args.ticker,
        args.nome,
        args.codigo_cvm,
        collect_cvm=args.collect_cvm,
        ri_urls=args.ri_url,
        collect_ri=not args.sem_ri_crawler,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
