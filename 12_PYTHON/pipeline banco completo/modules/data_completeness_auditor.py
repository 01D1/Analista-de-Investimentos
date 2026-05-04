"""
Auditoria de completude CVM/B3.

Este modulo diagnostica onde dados fundamentalistas podem se perder:
cadastro, CSV bruto CVM, normalizacao e escrita no Excel final.

A B3/yfinance cobre preco e liquidez. Demonstracoes financeiras (DRE, BP,
DFC, DFP e ITR) sao auditadas a partir da CVM.
"""

from __future__ import annotations

import copy
import json
import logging
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd

from config import settings as cfg
from config.empresa_loader import get_empresa, listar_tickers
from config.mapeamento_contas_geral import get_mapeamento
from modules._helpers import extrair_valores_df
from modules.coletor_cvm import ColetorCVM
from modules.normalizador import Normalizador

logger = logging.getLogger("pipeline.data_quality")

DOCS_ESPERADOS = ("dre", "bpa", "bpp", "dfc_mi", "dfc_md")
TIPOS_CVM = ("con", "ind")

CRITICAS = {
    "bank": {
        "dre": [
            "margem_financeira_bruta",
            "provisao_credito",
            "receita_servicos",
            "resultado_operacional",
            "lucro_liquido",
        ],
        "balanco": [
            "ativo_total",
            "passivo_total",
            "ativos_remuneraveis",
            "patrimonio_liquido",
            "pl_controladores",
        ],
        "dfc": [],
    },
    "general": {
        "dre": [
            "receita_liquida",
            "lucro_bruto",
            "ebit",
            "resultado_antes_ir",
            "lucro_liquido",
        ],
        "balanco": [
            "ativo_total",
            "passivo_total",
            "patrimonio_liquido",
            "pl_controladores",
        ],
        "dfc": [
            "fluxo_operacional",
            "fluxo_investimento",
            "fluxo_financiamento",
            "capex_total",
        ],
    },
}

LABELS_EXCEL = {
    "bank": {
        "dre": {
            "margem_financeira_bruta": ["Margem Financeira Bruta"],
            "provisao_credito": ["Despesa de Provisao Expandida"],
            "receita_servicos": ["Receitas de Prestacao de Servicos"],
            "resultado_operacional": ["RESULTADO OPERACIONAL"],
            "lucro_liquido": ["LUCRO LIQUIDO (CONTROLADORES)", "Lucro Liquido"],
        },
        "balanco": {
            "ativo_total": ["ATIVO TOTAL", "Ativo Total"],
            "ativos_remuneraveis": ["Ativos Remuneraveis"],
            "patrimonio_liquido": ["PATRIMONIO LIQUIDO", "Patrimonio Liquido"],
            "pl_controladores": ["PL Controladores", "Patrimonio Liquido"],
        },
    },
    "general": {
        "dre": {
            "receita_liquida": ["Receita Liquida", "Receita de Venda de Bens"],
            "lucro_bruto": ["Lucro Bruto", "RESULTADO BRUTO"],
            "ebit": ["EBIT"],
            "ebitda": ["EBITDA"],
            "resultado_antes_ir": ["Resultado Antes dos Tributos"],
            "lucro_liquido": ["Lucro Liquido", "LUCRO LIQUIDO"],
        },
        "balanco": {
            "ativo_total": ["Ativo Total", "ATIVO TOTAL"],
            "passivo_total": ["Passivo Total", "PASSIVO TOTAL"],
            "patrimonio_liquido": ["Patrimonio Liquido"],
            "pl_controladores": ["PL Controladores"],
            "divida_liquida": ["Divida Liquida"],
        },
        "dfc": {
            "fluxo_operacional": ["Fluxo Operacional"],
            "capex_total": ["CAPEX Total"],
        },
    },
}


def _eh_df_valido(df: Any) -> bool:
    return isinstance(df, pd.DataFrame) and not df.empty


def _safe_float(valor: Any) -> float | None:
    try:
        if pd.isna(valor):
            return None
        return float(valor)
    except Exception:
        return None


def _json_default(obj: Any) -> Any:
    if isinstance(obj, Path):
        return str(obj)
    if isinstance(obj, (pd.Timestamp, datetime)):
        return obj.isoformat()
    if hasattr(obj, "item"):
        try:
            return obj.item()
        except Exception:
            pass
    return str(obj)


def _doc_counts(dados_doc: dict[str, pd.DataFrame]) -> dict[str, int]:
    return {
        doc: int(len(dados_doc.get(doc))) if _eh_df_valido(dados_doc.get(doc)) else 0
        for doc in DOCS_ESPERADOS
    }


def _tem_docs_minimos(counts: dict[str, int]) -> bool:
    return counts.get("dre", 0) > 0 and counts.get("bpa", 0) > 0 and counts.get("bpp", 0) > 0


def _escolher_tipo(dfp_por_tipo: dict[str, dict[str, pd.DataFrame]]) -> str | None:
    counts_con = _doc_counts(dfp_por_tipo.get("con", {}))
    if _tem_docs_minimos(counts_con):
        return "con"

    counts_ind = _doc_counts(dfp_por_tipo.get("ind", {}))
    if _tem_docs_minimos(counts_ind):
        return "ind"

    placar = {}
    for tipo, dados in dfp_por_tipo.items():
        counts = _doc_counts(dados)
        placar[tipo] = (
            int(counts.get("dre", 0) > 0)
            + int(counts.get("bpa", 0) > 0)
            + int(counts.get("bpp", 0) > 0),
            sum(counts.values()),
        )
    melhor = max(placar.items(), key=lambda item: item[1], default=(None, (0, 0)))
    return melhor[0] if melhor[1] != (0, 0) else None


def _serie_por_linha(df: pd.DataFrame, linha: str) -> dict[int, float]:
    if df is None or df.empty or linha not in df.index:
        return {}
    serie = {}
    for ano, valor in df.loc[linha].items():
        try:
            serie[int(ano)] = float(valor)
        except Exception:
            continue
    return serie


def _anos_com_dado(df: pd.DataFrame, linhas: list[str] | None = None) -> list[int]:
    if df is None or df.empty:
        return []
    work = df
    if linhas:
        existentes = [l for l in linhas if l in df.index]
        if existentes:
            work = df.loc[existentes]
    anos = []
    for col in work.columns:
        vals = pd.to_numeric(work[col], errors="coerce").fillna(0)
        if (vals != 0).any():
            try:
                anos.append(int(col))
            except Exception:
                pass
    return sorted(anos)


def _linhas_zeradas(df: pd.DataFrame, linhas: list[str]) -> dict[str, list[int]]:
    resultado = {}
    if df is None or df.empty:
        return {linha: [] for linha in linhas}
    for linha in linhas:
        if linha not in df.index:
            resultado[linha] = []
            continue
        zeros = []
        for ano, valor in df.loc[linha].items():
            val = _safe_float(valor)
            if val is None or val == 0:
                try:
                    zeros.append(int(ano))
                except Exception:
                    pass
        resultado[linha] = zeros
    return resultado


def _aplicar_overrides(mapas: dict[str, dict], empresa: dict | None) -> dict[str, dict]:
    mapas = {
        "dre": copy.deepcopy(mapas.get("dre", {})),
        "ativo": copy.deepcopy(mapas.get("ativo", {})),
        "passivo": copy.deepcopy(mapas.get("passivo", {})),
        "dfc": copy.deepcopy(mapas.get("dfc", {})),
    }
    overrides = (empresa or {}).get("mapa_override", {})
    for conta, override in overrides.items():
        if conta in mapas["ativo"]:
            mapas["ativo"][conta].update(override)
        elif conta in mapas["passivo"]:
            mapas["passivo"][conta].update(override)
        elif conta in mapas["dre"]:
            mapas["dre"][conta].update(override)
        elif conta in mapas["dfc"]:
            mapas["dfc"][conta].update(override)
    return mapas


def _normalizar(
    dados_brutos: dict[int, dict[str, pd.DataFrame]],
    empresa: dict | None,
    tipo_empresa: str,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, dict[str, dict]]:
    mapas = _aplicar_overrides(get_mapeamento(tipo_empresa), empresa)
    divisor = (empresa or {}).get("divisor_cvm") or cfg.DIVISOR_VALORES
    norm = Normalizador(divisor=divisor)
    dre = norm.normalizar_dre(dados_brutos, mapas["dre"])
    bp = norm.normalizar_balanco(dados_brutos, mapas["ativo"], mapas["passivo"])
    dfc = pd.DataFrame()
    if tipo_empresa != "bank" and mapas.get("dfc"):
        dfc = norm.normalizar_dfc(dados_brutos, mapas["dfc"])
        dre = norm.inject_da(dre, bp)
    return dre, bp, dfc, mapas


def auditar_cadastro(ticker: str) -> dict[str, Any]:
    """Audita cadastro do ticker em config/empresas.yaml."""
    ticker = ticker.upper().strip()
    empresa = get_empresa(ticker)
    if not empresa:
        return {
            "ticker": ticker,
            "encontrado_empresas_yaml": False,
            "tem_codigo_cvm": False,
            "status": "critico",
            "problemas": ["Ticker nao encontrado em config/empresas.yaml."],
        }

    problemas = []
    if not empresa.get("codigo_cvm"):
        problemas.append("Empresa sem codigo_cvm cadastrado.")
    if not empresa.get("setor"):
        problemas.append("Empresa sem setor cadastrado.")
    if not empresa.get("tipo_empresa"):
        problemas.append("Tipo de empresa nao resolvido pelo setor.")

    return {
        "ticker": empresa.get("ticker", ticker),
        "nome": empresa.get("nome"),
        "encontrado_empresas_yaml": True,
        "tem_codigo_cvm": bool(empresa.get("codigo_cvm")),
        "codigo_cvm": empresa.get("codigo_cvm"),
        "setor": empresa.get("setor"),
        "tipo_empresa": empresa.get("tipo_empresa"),
        "divisor_cvm": empresa.get("divisor_cvm") or cfg.DIVISOR_VALORES,
        "aliases": empresa.get("aliases", []),
        "status": "ok" if not problemas else "atencao",
        "problemas": problemas,
    }


def _coletar_raw_cvm(
    ticker: str,
    anos: list[int],
    usar_cache: bool = True,
    incluir_itr: bool = True,
) -> tuple[dict[str, Any], dict[int, dict[str, pd.DataFrame]]]:
    empresa = get_empresa(ticker)
    codigo_cvm = (empresa or {}).get("codigo_cvm")
    coletor = ColetorCVM(cfg.CACHE_DIR, usar_cache=usar_cache)

    if not codigo_cvm:
        codigo_cvm = coletor.buscar_codigo_cvm(ticker)

    auditoria: dict[str, Any] = {
        "ticker": ticker.upper(),
        "codigo_cvm": codigo_cvm,
        "anos": {},
        "problemas": [],
    }
    dados_preferidos: dict[int, dict[str, pd.DataFrame]] = {}

    if not codigo_cvm:
        auditoria["problemas"].append("Codigo CVM ausente; raw CVM nao auditado.")
        return auditoria, dados_preferidos

    for ano in sorted({int(a) for a in anos}):
        ano_info = {
            "dfp": {},
            "itr": {},
            "tipo_preferido_dfp": None,
            "fonte_normalizacao": None,
            "problemas": [],
        }
        dfp_por_tipo: dict[str, dict[str, pd.DataFrame]] = {}

        for tipo in TIPOS_CVM:
            try:
                dados_dfp = coletor.baixar_dfp(str(codigo_cvm), ano, tipo=tipo)
            except Exception as exc:
                logger.warning("Erro ao baixar DFP %s %s %s: %s", ticker, ano, tipo, exc)
                dados_dfp = {}
                ano_info["problemas"].append(f"Erro ao baixar DFP {tipo}: {exc}")
            dfp_por_tipo[tipo] = dados_dfp
            counts = _doc_counts(dados_dfp)
            ano_info["dfp"][tipo] = {
                "docs_encontrados": [doc for doc, n in counts.items() if n > 0],
                "linhas_por_doc": counts,
                "tem_dre_bpa_bpp": _tem_docs_minimos(counts),
            }

        tipo_preferido = _escolher_tipo(dfp_por_tipo)
        ano_info["tipo_preferido_dfp"] = tipo_preferido
        if tipo_preferido:
            dados_preferidos[ano] = dfp_por_tipo[tipo_preferido]
            ano_info["fonte_normalizacao"] = f"dfp_{tipo_preferido}"
        else:
            ano_info["problemas"].append("Nenhum DFP con/ind encontrado para o ano.")

        if incluir_itr:
            for tipo in TIPOS_CVM:
                try:
                    dados_itr = coletor.baixar_itr(str(codigo_cvm), ano, trimestre=None, tipo=tipo)
                except Exception as exc:
                    logger.warning("Erro ao baixar ITR %s %s %s: %s", ticker, ano, tipo, exc)
                    dados_itr = {}
                    ano_info["problemas"].append(f"Erro ao baixar ITR {tipo}: {exc}")
                counts_itr = _doc_counts(dados_itr)
                ano_info["itr"][tipo] = {
                    "docs_encontrados": [doc for doc, n in counts_itr.items() if n > 0],
                    "linhas_por_doc": counts_itr,
                    "tem_dre_bpa_bpp": _tem_docs_minimos(counts_itr),
                }

            if not tipo_preferido:
                for tipo in TIPOS_CVM:
                    try:
                        dados_itr_ltm = coletor.anualizar_via_itr(
                            str(codigo_cvm),
                            ano,
                            dados_preferidos.get(ano - 1, {}),
                            tipo=tipo,
                        )
                    except Exception as exc:
                        logger.warning(
                            "Erro ao anualizar ITR %s %s %s: %s", ticker, ano, tipo, exc
                        )
                        dados_itr_ltm = {}
                    if dados_itr_ltm and any(_eh_df_valido(v) for v in dados_itr_ltm.values()):
                        dados_preferidos[ano] = dados_itr_ltm
                        ano_info["fonte_normalizacao"] = f"itr_ltm_{tipo}"
                        break

        auditoria["anos"][str(ano)] = ano_info

    return auditoria, dados_preferidos


def auditar_raw_cvm(ticker: str, anos: list[int]) -> dict[str, Any]:
    """Audita disponibilidade bruta DFP/ITR na CVM por ano e por con/ind."""
    auditoria, _ = _coletar_raw_cvm(ticker, anos, usar_cache=True, incluir_itr=True)
    return auditoria


def auditar_b3_cache(ticker: str) -> dict[str, Any]:
    """Audita evidencia local de preco/volume B3 no cache de mercado."""
    ticker = ticker.upper().strip()
    padrao = f"precos_{ticker}.SA_*.parquet"
    arquivos = sorted(cfg.CACHE_DIR.glob(padrao), reverse=True)
    resultado: dict[str, Any] = {
        "ticker": ticker,
        "fonte": "cache/yfinance",
        "arquivos_encontrados": [str(p) for p in arquivos[:5]],
        "auditado": bool(arquivos),
        "problemas": [],
    }
    if not arquivos:
        resultado["problemas"].append("Nenhum cache de precos B3/yfinance encontrado.")
        return resultado

    path = arquivos[0]
    resultado["arquivo_usado"] = str(path)
    try:
        df = pd.read_parquet(path)
    except Exception as exc:
        resultado["problemas"].append(f"Nao foi possivel ler parquet de precos: {exc}")
        return resultado

    resultado["linhas"] = int(len(df))
    if df.empty:
        resultado["problemas"].append("Arquivo de precos vazio.")
        return resultado

    idx = df.index
    try:
        resultado["data_inicial"] = str(idx.min().date())
        resultado["data_final"] = str(idx.max().date())
    except Exception:
        pass

    col_volume = next((c for c in df.columns if str(c).lower() == "volume"), None)
    if col_volume:
        volume = pd.to_numeric(df[col_volume], errors="coerce").fillna(0)
        resultado["dias_com_volume"] = int((volume > 0).sum())
        resultado["volume_medio"] = float(volume.mean()) if not volume.empty else 0.0
        if resultado["dias_com_volume"] == 0:
            resultado["problemas"].append("Cache de precos sem volume negociado.")
    else:
        resultado["problemas"].append("Coluna de volume nao encontrada no cache de precos.")

    return resultado


def _auditar_mapeamento(
    dados_brutos: dict[int, dict[str, pd.DataFrame]],
    normalizado: pd.DataFrame,
    mapa: dict[str, dict],
    doc_tipo: str,
    divisor: float,
) -> list[dict[str, Any]]:
    problemas = []
    if not mapa:
        return problemas

    for conta, config in mapa.items():
        codigos = config.get("codigos", [])
        if config.get("calculado") or not codigos:
            continue
        for ano, docs in sorted(dados_brutos.items()):
            df = docs.get(doc_tipo)
            if not _eh_df_valido(df):
                continue
            raw = extrair_valores_df(df, codigos, periodo=f"{ano}-12-31")
            raw_norm = round(raw * config.get("sinal", 1) / divisor, 3)
            norm_val = None
            if normalizado is not None and not normalizado.empty:
                if conta in normalizado.index and ano in normalizado.columns:
                    norm_val = _safe_float(normalizado.loc[conta, ano])
            if raw_norm != 0 and (norm_val is None or norm_val == 0):
                problemas.append(
                    {
                        "ano": int(ano),
                        "conta": conta,
                        "doc": doc_tipo,
                        "codigos": codigos,
                        "valor_raw_normalizado": raw_norm,
                        "valor_dataframe": norm_val,
                    }
                )
    return problemas


def auditar_normalizacao(
    ticker: str,
    dre_hist: pd.DataFrame,
    bp_hist: pd.DataFrame,
    dfc_hist: pd.DataFrame | None = None,
    dados_brutos: dict[int, dict[str, pd.DataFrame]] | None = None,
    mapas: dict[str, dict] | None = None,
) -> dict[str, Any]:
    """Audita DataFrames normalizados e zeros suspeitos contra o raw CVM."""
    empresa = get_empresa(ticker)
    tipo_empresa = (empresa or {}).get("tipo_empresa", "general")
    criticas = CRITICAS.get(tipo_empresa, CRITICAS["general"])
    divisor = (empresa or {}).get("divisor_cvm") or cfg.DIVISOR_VALORES
    if mapas is None:
        mapas = _aplicar_overrides(get_mapeamento(tipo_empresa), empresa)
    if dfc_hist is None:
        dfc_hist = pd.DataFrame()

    resultado: dict[str, Any] = {
        "tipo_empresa": tipo_empresa,
        "anos_com_dre": _anos_com_dado(dre_hist, criticas["dre"]),
        "anos_com_bp": _anos_com_dado(bp_hist, criticas["balanco"]),
        "anos_com_dfc": _anos_com_dado(dfc_hist, criticas["dfc"]) if criticas["dfc"] else [],
        "linhas_criticas_zeradas": {
            "dre": _linhas_zeradas(dre_hist, criticas["dre"]),
            "balanco": _linhas_zeradas(bp_hist, criticas["balanco"]),
            "dfc": _linhas_zeradas(dfc_hist, criticas["dfc"]) if criticas["dfc"] else {},
        },
        "zeros_suspeitos_pos_normalizacao": [],
        "conciliacao_balanco": {},
        "problemas": [],
    }

    if dados_brutos:
        resultado["zeros_suspeitos_pos_normalizacao"].extend(
            _auditar_mapeamento(dados_brutos, dre_hist, mapas.get("dre", {}), "dre", divisor)
        )
        resultado["zeros_suspeitos_pos_normalizacao"].extend(
            _auditar_mapeamento(dados_brutos, bp_hist, mapas.get("ativo", {}), "bpa", divisor)
        )
        resultado["zeros_suspeitos_pos_normalizacao"].extend(
            _auditar_mapeamento(dados_brutos, bp_hist, mapas.get("passivo", {}), "bpp", divisor)
        )
        if mapas.get("dfc") and dfc_hist is not None and not dfc_hist.empty:
            resultado["zeros_suspeitos_pos_normalizacao"].extend(
                _auditar_mapeamento(
                    dados_brutos, dfc_hist, mapas.get("dfc", {}), "dfc_mi", divisor
                )
            )

    ativo = _serie_por_linha(bp_hist, "ativo_total")
    passivo = _serie_por_linha(bp_hist, "passivo_total")
    conciliacao = {}
    for ano in sorted(set(ativo) | set(passivo)):
        a = ativo.get(ano)
        p = passivo.get(ano)
        diff_pct = None
        if a not in (None, 0) and p is not None:
            diff_pct = abs(a - p) / abs(a)
        conciliacao[ano] = {"ativo_total": a, "passivo_total": p, "diff_pct": diff_pct}
    resultado["conciliacao_balanco"] = conciliacao

    if len(resultado["anos_com_dre"]) < 5:
        resultado["problemas"].append("DRE com menos de 5 anos historicos nao zerados.")
    if len(resultado["anos_com_bp"]) < 5:
        resultado["problemas"].append("BP com menos de 5 anos historicos nao zerados.")
    if tipo_empresa != "bank" and len(resultado["anos_com_dfc"]) < 5:
        resultado["problemas"].append("DFC com menos de 5 anos historicos nao zerados.")
    if resultado["zeros_suspeitos_pos_normalizacao"]:
        resultado["problemas"].append("Ha contas com valor raw nao zero e DataFrame zerado.")

    return resultado


def _mapear_colunas_anos(ws) -> dict[int, int]:
    mapa = {}
    for row in ws.iter_rows(min_row=1, max_row=8, values_only=True):
        for idx, val in enumerate(row, start=1):
            if val is None:
                continue
            try:
                ano = int(str(val).strip())
            except Exception:
                continue
            if 2010 <= ano <= 2045:
                mapa.setdefault(ano, col)
    return mapa


def _normalizar_label(valor: Any) -> str:
    if valor is None:
        return ""
    txt = str(valor).strip().lower()
    troca = str.maketrans(
        "áàãâäéèêëíìîïóòõôöúùûüçÁÀÃÂÄÉÈÊËÍÌÎÏÓÒÕÔÖÚÙÛÜÇ",
        "aaaaaeeeeiiiiooooouuuucAAAAAEEEEIIIIOOOOOUUUUC",
    )
    return txt.translate(troca)


def _procurar_linha(ws, labels: list[str]) -> int | None:
    labels_norm = [_normalizar_label(label) for label in labels]
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        for idx in (0, 3):
            valor = row[idx] if idx < len(row) else None
            valor_norm = _normalizar_label(valor)
            if not valor_norm:
                continue
            for label in labels_norm:
                if label and label in valor_norm:
                    return row_idx
    return None


def _indexar_aba_excel(ws) -> dict[str, Any]:
    """Indexa uma aba em uma passada para evitar leituras read_only muito lentas."""
    col_mapa: dict[int, int] = {}
    linhas_por_label: dict[str, int] = {}
    valores_por_linha: dict[int, tuple] = {}

    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row_idx <= 8:
            for col_idx, val in enumerate(row, start=1):
                if val is None:
                    continue
                try:
                    ano = int(str(val).strip())
                except Exception:
                    continue
                if 2010 <= ano <= 2045:
                    col_mapa.setdefault(ano, col_idx)

        tem_label = False
        for idx in (0, 3):
            valor = row[idx] if idx < len(row) else None
            label_norm = _normalizar_label(valor)
            if label_norm:
                linhas_por_label[label_norm] = row_idx
                tem_label = True

        if tem_label:
            valores_por_linha[row_idx] = tuple(row)

    return {
        "nome": ws.title,
        "col_mapa": col_mapa,
        "linhas_por_label": linhas_por_label,
        "valores_por_linha": valores_por_linha,
    }


def _match_linha_indexada(sheet_idx: dict[str, Any], labels: list[str]) -> int | None:
    labels_norm = [_normalizar_label(label) for label in labels]

    for label in labels_norm:
        row_idx = sheet_idx["linhas_por_label"].get(label)
        if row_idx:
            return row_idx

    for label in labels_norm:
        for valor_norm, row_idx in sheet_idx["linhas_por_label"].items():
            if not label:
                continue
            if "/" in valor_norm or "%" in valor_norm:
                continue
            if valor_norm.startswith(label):
                return row_idx
    return None


def _valor_indexado(sheet_idx: dict[str, Any], row_idx: int, col_idx: int) -> Any:
    row = sheet_idx["valores_por_linha"].get(row_idx, ())
    pos = col_idx - 1
    return row[pos] if pos < len(row) else None


def auditar_excel(
    ticker: str,
    excel_path: str | Path | None,
    dre_hist: pd.DataFrame,
    bp_hist: pd.DataFrame,
    dfc_hist: pd.DataFrame | None = None,
) -> dict[str, Any]:
    """Compara dados normalizados contra o Excel final, quando disponivel."""
    if not excel_path:
        return {"auditado": False, "motivo": "Excel final nao informado."}

    path = Path(excel_path)
    if not path.exists():
        return {"auditado": False, "excel_path": str(path), "motivo": "Arquivo nao encontrado."}

    try:
        from openpyxl import load_workbook
    except Exception as exc:
        return {"auditado": False, "excel_path": str(path), "motivo": f"openpyxl indisponivel: {exc}"}

    empresa = get_empresa(ticker)
    tipo_empresa = (empresa or {}).get("tipo_empresa", "general")
    labels = LABELS_EXCEL.get(tipo_empresa, LABELS_EXCEL["general"])
    dfc_excel = dfc_hist if dfc_hist is not None else pd.DataFrame()
    dfs = {"dre": dre_hist, "balanco": bp_hist, "dfc": dfc_excel}

    resultado = {
        "auditado": True,
        "excel_path": str(path),
        "abas": {},
        "contas_nao_encontradas": [],
        "divergencias": [],
        "anos_ausentes": [],
    }

    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        anos_normalizados = sorted(
            {
                int(c)
                for df in dfs.values()
                if df is not None and not df.empty
                for c in df.columns
                if str(c).isdigit()
            }
        )
        sheet_indexes = [_indexar_aba_excel(ws) for ws in wb.worksheets]

        anos_excel = set()
        for sheet_idx in sheet_indexes:
            col_mapa = sheet_idx["col_mapa"]
            if col_mapa:
                resultado["abas"][sheet_idx["nome"]] = {"anos": sorted(col_mapa)}
                anos_excel.update(col_mapa)
        resultado["anos_ausentes"] = [a for a in anos_normalizados if a not in anos_excel]

        for grupo, contas in labels.items():
            df = dfs.get(grupo)
            if df is None or df.empty:
                continue
            for conta, possiveis_labels in contas.items():
                if conta not in df.index:
                    continue
                achou = False
                for sheet_idx in sheet_indexes:
                    col_mapa = sheet_idx["col_mapa"]
                    if not col_mapa:
                        continue
                    row = _match_linha_indexada(sheet_idx, possiveis_labels)
                    if not row:
                        continue
                    achou = True
                    for ano in anos_normalizados:
                        col = col_mapa.get(ano)
                        if not col or ano not in df.columns:
                            continue
                        excel_val = _valor_indexado(sheet_idx, row, col)
                        if isinstance(excel_val, str) and excel_val.startswith("="):
                            continue
                        excel_num = _safe_float(excel_val)
                        df_num = _safe_float(df.loc[conta, ano])
                        if excel_num is None or df_num is None:
                            continue
                        if abs(excel_num - df_num) > max(1.0, abs(df_num) * 0.005):
                            resultado["divergencias"].append(
                                {
                                    "aba": sheet_idx["nome"],
                                    "conta": conta,
                                    "ano": ano,
                                    "dataframe": df_num,
                                    "excel": excel_num,
                                }
                            )
                    break
                if not achou:
                    resultado["contas_nao_encontradas"].append({"grupo": grupo, "conta": conta})
    finally:
        wb.close()

    return resultado


def _ultimo_excel_ticker(ticker: str) -> Path | None:
    ticker_up = ticker.upper()
    candidatos = [
        p
        for p in cfg.OUTPUT_DIR.glob(f"Valuation_{ticker_up}_*.xlsx")
        if "test" not in p.name.lower() and "summary" not in p.name.lower()
    ]
    canonico = cfg.OUTPUT_DIR / "valuations" / ticker_up / f"Valuation_{ticker_up}.xlsx"
    if canonico.exists():
        candidatos.append(canonico)
    if not candidatos:
        candidatos = list(cfg.OUTPUT_DIR.glob(f"Valuation_{ticker_up}_*.xlsx"))
    candidatos = sorted(candidatos, key=lambda p: p.stat().st_mtime, reverse=True)
    return candidatos[0] if candidatos else None


def _score_auditoria(auditoria: dict[str, Any]) -> tuple[int, list[str], list[str]]:
    warnings: list[str] = []
    critical: list[str] = []
    score = 100

    cadastro = auditoria.get("cadastro", {})
    if not cadastro.get("encontrado_empresas_yaml"):
        critical.append("Ticker nao cadastrado em empresas.yaml.")
    elif not cadastro.get("tem_codigo_cvm"):
        critical.append("Ticker sem codigo_cvm no cadastro.")

    b3 = auditoria.get("b3", {})
    for problema in b3.get("problemas", []):
        warnings.append(f"B3/cache: {problema}")

    raw = auditoria.get("raw_cvm", {})
    for ano, info in raw.get("anos", {}).items():
        tipo = info.get("tipo_preferido_dfp")
        if not tipo:
            fonte_norm = str(info.get("fonte_normalizacao") or "")
            if fonte_norm.startswith("itr_ltm"):
                warnings.append(f"{ano}: DFP ausente; normalizacao usou fallback {fonte_norm}.")
                continue
            critical.append(f"{ano}: nenhum DFP con/ind encontrado.")
            continue
        counts = info.get("dfp", {}).get(tipo, {}).get("linhas_por_doc", {})
        if counts.get("dre", 0) == 0:
            critical.append(f"{ano}: DRE bruta ausente no DFP {tipo}.")
        if counts.get("bpa", 0) == 0 or counts.get("bpp", 0) == 0:
            critical.append(f"{ano}: BPA/BPP bruto ausente no DFP {tipo}.")
        if counts.get("dfc_mi", 0) == 0 and counts.get("dfc_md", 0) == 0:
            warnings.append(f"{ano}: DFC bruta ausente no DFP {tipo}.")

    normalizacao = auditoria.get("normalizacao", {})
    for problema in normalizacao.get("problemas", []):
        critical.append(problema)
    if normalizacao.get("zeros_suspeitos_pos_normalizacao"):
        critical.append("Zeros suspeitos detectados na normalizacao.")

    excel = auditoria.get("excel", {})
    if excel.get("auditado"):
        if excel.get("anos_ausentes"):
            critical.append("Excel sem colunas para anos historicos normalizados.")
        if excel.get("divergencias"):
            critical.append("Divergencias entre DataFrame normalizado e Excel final.")
        if excel.get("contas_nao_encontradas"):
            warnings.append("Algumas contas criticas nao foram localizadas no Excel.")

    score -= len(critical) * 18
    score -= len(warnings) * 5
    return max(score, 0), warnings, critical


def gerar_relatorio_completude(
    ticker: str,
    auditoria: dict[str, Any],
    output_dir: str | Path | None = None,
) -> Path:
    """Salva relatorio Markdown e JSON da auditoria."""
    out_dir = Path(output_dir) if output_dir else cfg.OUTPUT_DIR / "data_quality"
    out_dir.mkdir(parents=True, exist_ok=True)
    data_str = datetime.now().strftime("%Y%m%d")
    base = out_dir / f"audit_{ticker.upper()}_{data_str}"

    score, warnings, critical = _score_auditoria(auditoria)
    auditoria["score"] = score
    auditoria["warnings"] = warnings
    auditoria["critical"] = critical

    json_path = base.with_suffix(".json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(auditoria, f, ensure_ascii=False, indent=2, default=_json_default)

    md_path = base.with_suffix(".md")
    cadastro = auditoria.get("cadastro", {})
    b3 = auditoria.get("b3", {})
    normalizacao = auditoria.get("normalizacao", {})
    excel = auditoria.get("excel", {})

    linhas = [
        f"# Auditoria de Completude CVM/B3 - {ticker.upper()}",
        "",
        f"Data: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"Score: {score}/100",
        "",
        "## Cadastro",
        "",
        f"- Nome: {cadastro.get('nome', '-')}",
        f"- Codigo CVM: {cadastro.get('codigo_cvm') or 'AUSENTE'}",
        f"- Setor: {cadastro.get('setor', '-')}",
        f"- Tipo: {cadastro.get('tipo_empresa', '-')}",
        "",
        "## B3 / Mercado",
        "",
        f"- Fonte: {b3.get('fonte', '-')}",
        f"- Arquivo usado: {b3.get('arquivo_usado', '-')}",
        f"- Linhas: {b3.get('linhas', 0)}",
        f"- Periodo: {b3.get('data_inicial', '-')} a {b3.get('data_final', '-')}",
        f"- Dias com volume: {b3.get('dias_com_volume', '-')}",
        "",
        "## Raw CVM",
        "",
        "| Ano | Tipo DFP usado | DRE | BPA | BPP | DFC_MI | DFC_MD |",
        "| --- | --- | ---: | ---: | ---: | ---: | ---: |",
    ]
    for ano, info in auditoria.get("raw_cvm", {}).get("anos", {}).items():
        tipo = info.get("tipo_preferido_dfp") or "-"
        counts = info.get("dfp", {}).get(tipo, {}).get("linhas_por_doc", {}) if tipo != "-" else {}
        linhas.append(
            f"| {ano} | {tipo} | {counts.get('dre', 0)} | {counts.get('bpa', 0)} | "
            f"{counts.get('bpp', 0)} | {counts.get('dfc_mi', 0)} | {counts.get('dfc_md', 0)} |"
        )

    linhas.extend(
        [
            "",
            "## Normalizacao",
            "",
            f"- Anos com DRE: {normalizacao.get('anos_com_dre', [])}",
            f"- Anos com BP: {normalizacao.get('anos_com_bp', [])}",
            f"- Anos com DFC: {normalizacao.get('anos_com_dfc', [])}",
            f"- Zeros suspeitos: {len(normalizacao.get('zeros_suspeitos_pos_normalizacao', []))}",
            "",
            "## Excel",
            "",
        ]
    )
    if excel.get("auditado"):
        linhas.extend(
            [
                f"- Arquivo: {excel.get('excel_path')}",
                f"- Anos ausentes: {excel.get('anos_ausentes', [])}",
                f"- Divergencias: {len(excel.get('divergencias', []))}",
                f"- Contas nao encontradas: {len(excel.get('contas_nao_encontradas', []))}",
            ]
        )
    else:
        linhas.append(f"- Nao auditado: {excel.get('motivo', '-')}")

    linhas.extend(["", "## Falhas criticas", ""])
    linhas.extend([f"- {item}" for item in critical] or ["- Nenhuma."])
    linhas.extend(["", "## Avisos", ""])
    linhas.extend([f"- {item}" for item in warnings] or ["- Nenhum."])

    with open(md_path, "w", encoding="utf-8") as f:
        f.write("\n".join(linhas) + "\n")

    auditoria["arquivos"] = {"json": str(json_path), "markdown": str(md_path)}
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(auditoria, f, ensure_ascii=False, indent=2, default=_json_default)

    return md_path


def executar_auditoria_completude(
    ticker: str,
    anos: list[int] | None = None,
    excel_path: str | Path | None = None,
    usar_cache: bool = True,
    output_dir: str | Path | None = None,
    incluir_itr: bool = True,
) -> dict[str, Any]:
    """Executa a auditoria completa de um ticker e salva relatorios."""
    ticker = ticker.upper().strip()
    anos = anos or cfg.ANOS_HISTORICOS

    cadastro = auditar_cadastro(ticker)
    empresa = get_empresa(ticker)
    tipo_empresa = (empresa or {}).get("tipo_empresa", "general")

    raw_cvm, dados_brutos = _coletar_raw_cvm(
        ticker=ticker,
        anos=anos,
        usar_cache=usar_cache,
        incluir_itr=incluir_itr,
    )

    dre_hist, bp_hist, dfc_hist, mapas = _normalizar(dados_brutos, empresa, tipo_empresa)

    normalizacao = auditar_normalizacao(
        ticker=ticker,
        dre_hist=dre_hist,
        bp_hist=bp_hist,
        dfc_hist=dfc_hist,
        dados_brutos=dados_brutos,
        mapas=mapas,
    )

    if excel_path is None:
        excel_path = _ultimo_excel_ticker(ticker)
    excel = auditar_excel(ticker, excel_path, dre_hist, bp_hist, dfc_hist)
    b3 = auditar_b3_cache(ticker)

    auditoria = {
        "ticker": ticker,
        "cadastro": cadastro,
        "b3": b3,
        "raw_cvm": raw_cvm,
        "normalizacao": normalizacao,
        "excel": excel,
    }
    md_path = gerar_relatorio_completude(ticker, auditoria, output_dir=output_dir)
    auditoria["arquivos"] = {
        "markdown": str(md_path),
        "json": str(Path(md_path).with_suffix(".json")),
    }
    return auditoria


def gerar_indice_csv(auditorias: list[dict[str, Any]], output_dir: str | Path | None = None) -> Path:
    """Gera CSV consolidado das auditorias executadas."""
    out_dir = Path(output_dir) if output_dir else cfg.OUTPUT_DIR / "data_quality"
    out_dir.mkdir(parents=True, exist_ok=True)
    data_str = datetime.now().strftime("%Y%m%d")
    csv_path = out_dir / f"audit_all_{data_str}.csv"

    linhas = []
    for audit in auditorias:
        cad = audit.get("cadastro", {})
        norm = audit.get("normalizacao", {})
        excel = audit.get("excel", {})
        b3 = audit.get("b3", {})
        linhas.append(
            {
                "ticker": audit.get("ticker"),
                "nome": cad.get("nome"),
                "codigo_cvm": cad.get("codigo_cvm"),
                "setor": cad.get("setor"),
                "tipo_empresa": cad.get("tipo_empresa"),
                "score": audit.get("score"),
                "critical": len(audit.get("critical", [])),
                "warnings": len(audit.get("warnings", [])),
                "anos_dre": len(norm.get("anos_com_dre", [])),
                "anos_bp": len(norm.get("anos_com_bp", [])),
                "anos_dfc": len(norm.get("anos_com_dfc", [])),
                "zeros_suspeitos": len(norm.get("zeros_suspeitos_pos_normalizacao", [])),
                "b3_cache_auditado": b3.get("auditado", False),
                "b3_linhas": b3.get("linhas"),
                "b3_dias_com_volume": b3.get("dias_com_volume"),
                "excel_auditado": excel.get("auditado", False),
                "excel_divergencias": len(excel.get("divergencias", [])),
                "excel_anos_ausentes": len(excel.get("anos_ausentes", [])),
                "relatorio": audit.get("arquivos", {}).get("markdown"),
            }
        )

    df_novo = pd.DataFrame(linhas)
    if csv_path.exists():
        try:
            df_existente = pd.read_csv(csv_path)
            if not df_existente.empty:
                df_final = pd.concat([df_existente, df_novo], ignore_index=True)
                df_final = df_final.drop_duplicates(subset=["ticker"], keep="last")
            else:
                df_final = df_novo
        except Exception:
            df_final = df_novo
    else:
        df_final = df_novo

    if not df_final.empty and "ticker" in df_final.columns:
        df_final = df_final.sort_values("ticker")

    df_final.to_csv(csv_path, index=False, encoding="utf-8-sig")
    return csv_path


def listar_tickers_auditoria(all_tickers: bool, ticker: str | None = None) -> list[str]:
    """Helper para CLI."""
    if all_tickers:
        return listar_tickers()
    return [ticker.upper().strip()] if ticker else []
