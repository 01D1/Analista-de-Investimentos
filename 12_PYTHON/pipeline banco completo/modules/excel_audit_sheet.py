"""
excel_audit_sheet.py

Aba executiva de auditoria para os workbooks de valuation. Ela concentra
status do modelo, qualidade de dados, classe da acao, fontes e premissas para
o analista enxergar rapidamente se a planilha e confiavel ou preliminar.
"""

from __future__ import annotations

from datetime import datetime
from typing import Any

from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


HEADER = "1F4E79"
SUBHEADER = "D9EAF7"
OK = "E2F0D9"
WARN = "FFF2CC"
BAD = "FCE4D6"
WHITE = "FFFFFF"
BORDER = "D9D9D9"


def _fill(color: str) -> PatternFill:
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


def _border() -> Border:
    side = Side(style="thin", color=BORDER)
    return Border(left=side, right=side, top=side, bottom=side)


def _safe(value: Any, dash: str = "-") -> Any:
    if value is None:
        return dash
    if isinstance(value, float):
        return round(value, 6)
    return value


def _pct(value: Any) -> Any:
    return value if isinstance(value, (int, float)) else _safe(value)


def _status_fill(status: str) -> str:
    status = str(status or "").upper()
    if "BLOQUE" in status:
        return BAD
    if "PRELIMINAR" in status or "REVIS" in status:
        return WARN
    return OK


def adicionar_aba_auditoria(wb, dados: dict, nome_empresa: str, ticker: str) -> None:
    """Cria/atualiza a aba Status & Fontes."""
    sheet_name = "Status & Fontes"
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name, 0)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A7"

    valuation = dados.get("valuation", {}) or {}
    mercado = dados.get("mercado", {}) or {}
    metodologia = dados.get("metodologia", {}) or {}
    prem = metodologia.get("premissas", {}) or {}
    prem_ef = metodologia.get("premissas_efetivas", {}) or {}
    qualidade = dados.get("qualidade", {}) or {}
    qualitativo = dados.get("qualitativo", {}) or {}

    for col, width in {
        "A": 26, "B": 28, "C": 18, "D": 20, "E": 20, "F": 50,
    }.items():
        ws.column_dimensions[col].width = width

    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = f"Status & Fontes - {nome_empresa} ({ticker})"
    c.fill = _fill(HEADER)
    c.font = Font(color=WHITE, bold=True, size=14)
    c.alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 28

    status = valuation.get("status_valuation") or qualidade.get("status_valuation") or "CONFIAVEL"
    ws.merge_cells("A3:B4")
    ws["A3"] = "STATUS DO VALUATION"
    ws["A3"].font = Font(bold=True, color=WHITE)
    ws["A3"].fill = _fill(HEADER)
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("C3:F4")
    ws["C3"] = status
    ws["C3"].font = Font(bold=True, size=16)
    ws["C3"].fill = _fill(_status_fill(status))
    ws["C3"].alignment = Alignment(horizontal="center", vertical="center")
    for row in ws["A3:F4"]:
        for cell in row:
            cell.border = _border()

    def section(row: int, title: str):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        cell = ws.cell(row=row, column=1, value=title)
        cell.fill = _fill(HEADER)
        cell.font = Font(color=WHITE, bold=True)
        cell.border = _border()
        return row + 1

    def row_values(row: int, values: list[Any], fills: list[str] | None = None):
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=_safe(value))
            cell.border = _border()
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if col == 1:
                cell.font = Font(bold=True)
                cell.fill = _fill(SUBHEADER)
            if fills and col <= len(fills) and fills[col - 1]:
                cell.fill = _fill(fills[col - 1])

    row = section(6, "Resumo executivo")
    resumo = [
        ("Classe principal", valuation.get("classe_principal") or mercado.get("classe_principal")),
        ("Tipo de acao", valuation.get("tipo_acao") or metodologia.get("tipo_acao")),
        ("Preco justo principal", valuation.get("preco_justo_principal")),
        ("Upside principal", valuation.get("upside_principal")),
        ("TIR principal", valuation.get("tir_principal")),
        ("Quality score pre-valuation", qualidade.get("score")),
        ("Alertas pre-valuation", len(qualidade.get("warnings", []) or [])),
        ("Criticos pre-valuation", len(qualidade.get("critical", []) or [])),
        ("Score qualitativo", qualitativo.get("overall_score")),
    ]
    for label, value in resumo:
        row_values(row, [label, value, "", "", "", ""])
        if "Upside" in label or "TIR" in label:
            ws.cell(row=row, column=2).number_format = "0.0%;(0.0%);-"
        elif "Preco" in label:
            ws.cell(row=row, column=2).number_format = 'R$ #,##0.00;(R$ #,##0.00);"-"'
        row += 1

    row += 1
    row = section(row, "Fontes dos dados")
    fontes = [
        ("Demonstrativos historicos", "CVM DFP/ITR", "DRE, BP e DFC normalizados", "automatico", "", ""),
        ("Mercado", "B3/yfinance/cache ou CLI", "cotacao, beta estatistico, volume/liquidez quando disponivel", "automatico/input", "", ""),
        ("Macro", "BCB/Focus/cache/fallback", "Selic, DI, IPCA, CDS/TJLP quando aplicavel", "automatico/fallback", "", ""),
        ("RI/CVM qualitativo", "CVM IPE + crawler RI", "releases, fatos relevantes, comunicados e eventos", "automatico/configurado", "", ""),
    ]
    row_values(row, ["Item", "Fonte", "Uso no modelo", "Tipo", "Periodo/as-of", "Notas"], [SUBHEADER] * 6)
    for col in range(1, 7):
        ws.cell(row=row, column=col).font = Font(bold=True)
    row += 1
    for vals in fontes:
        row_values(row, list(vals))
        row += 1

    row += 1
    row = section(row, "Premissas-chave e origem")
    prem_rows = [
        ("g perpetuidade", prem_ef.get("g_perpetuidade") or valuation.get("g_perpetuidade"), "empresas.yaml/settings/CLI", "subjetiva"),
        ("Beta usado", prem_ef.get("beta_usado") or mercado.get("beta_usar"), prem_ef.get("beta_fonte") or mercado.get("beta_fonte"), "subjetiva/mercado"),
        ("Payout", prem_ef.get("payout") or prem.get("payout"), "empresas.yaml/setor", "subjetiva"),
        ("Margem EBITDA alvo", prem_ef.get("margem_ebitda_alvo") or prem.get("margem_ebitda_alvo"), "empresas.yaml/setor", "subjetiva"),
        ("CAPEX / Receita", prem_ef.get("capex_pct_receita") or prem.get("capex_pct_receita"), "empresas.yaml/setor", "subjetiva"),
        ("NCG / Receita", prem_ef.get("ncg_pct_receita") or prem.get("ncg_pct_receita"), "empresas.yaml/setor", "subjetiva"),
        ("Custo divida spread", prem_ef.get("custo_divida_spread") or prem.get("custo_divida_spread"), "empresas.yaml/settings", "subjetiva"),
    ]
    row_values(row, ["Premissa", "Valor", "Fonte", "Categoria", "", ""], [SUBHEADER] * 6)
    for col in range(1, 7):
        ws.cell(row=row, column=col).font = Font(bold=True)
    row += 1
    for label, value, fonte, categoria in prem_rows:
        if value is None:
            continue
        row_values(row, [label, _pct(value), fonte, categoria, "", ""])
        if isinstance(value, (int, float)) and ("Beta" not in label):
            ws.cell(row=row, column=2).number_format = "0.0%;(0.0%);-"
        row += 1

    row += 1
    row = section(row, "Checks rapidos")
    checks = [
        ("Sem criticos pre-valuation", len(qualidade.get("critical", []) or []) == 0, "Se NAO, bloquear ou revisar antes de usar."),
        ("Valuation nao bloqueado", not qualidade.get("bloqueia_valuation", False), "Se NAO, resultado nao deve ser usado."),
        ("Classe principal definida", bool(valuation.get("classe_principal")), "Evita confundir ON/PN/UNIT."),
        ("Preco principal preenchido", valuation.get("preco_justo_principal") not in {None, 0}, "Valuation por acao precisa estar visivel."),
    ]
    row_values(row, ["Check", "Status", "Nota", "", "", ""], [SUBHEADER] * 6)
    for col in range(1, 7):
        ws.cell(row=row, column=col).font = Font(bold=True)
    row += 1
    for label, ok, note in checks:
        row_values(row, [label, "OK" if ok else "REVISAR", note, "", "", ""])
        ws.cell(row=row, column=2).fill = _fill(OK if ok else WARN)
        ws.cell(row=row, column=2).font = Font(bold=True)
        row += 1

    row += 1
    row_values(row, ["Gerado em", datetime.now().strftime("%Y-%m-%d %H:%M:%S"), "", "", "", ""])

    for row_cells in ws.iter_rows():
        for cell in row_cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for idx in range(1, row + 1):
        ws.row_dimensions[idx].height = 20
    ws.auto_filter.ref = f"A6:F{max(row, 6)}"

