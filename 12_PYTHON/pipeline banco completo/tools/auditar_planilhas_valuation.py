from __future__ import annotations

import csv
import math
import re
import unicodedata
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass, asdict
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET
from typing import Any

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover
    raise SystemExit("openpyxl nao esta disponivel neste ambiente") from exc


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = PROJECT_ROOT / "outputs" / "revisao_valuation"
REPORT_PATH = PROJECT_ROOT / "RELATORIO_REVISAO_VALUATION.md"
CURRENT_YEAR = datetime.now().year
FORMULA_ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!")
MAX_SCAN_ROWS = 500
MAX_SCAN_COLS = 80


@dataclass
class Finding:
    arquivo: str
    aba: str
    celula: str
    gravidade: str
    categoria: str
    problema: str
    explicacao: str
    impacto: str
    sugestao: str


def norm(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[^a-z0-9%]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def as_float(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if math.isnan(value) or math.isinf(value):
            return None
        return float(value)
    if isinstance(value, str):
        text = value.strip()
        if not text or text in {"-", "--"}:
            return None
        if text.startswith("="):
            return None
        text = text.replace("R$", "").replace("%", "").replace(" ", "")
        if "," in text and "." in text:
            text = text.replace(".", "").replace(",", ".")
        elif "," in text:
            text = text.replace(",", ".")
        try:
            return float(text)
        except ValueError:
            return None
    return None


def rel_addr(row: int, col: int) -> str:
    return f"{openpyxl.utils.get_column_letter(col)}{row}"


def rel_file(path: Path) -> str:
    try:
        return str(path.relative_to(PROJECT_ROOT))
    except ValueError:
        return str(path)


def discover_files() -> list[Path]:
    exts = {".xlsx", ".xlsm", ".xls", ".csv"}
    files = [p for p in PROJECT_ROOT.rglob("*") if p.is_file() and p.suffix.lower() in exts]
    valuation_files = []
    for path in files:
        p = str(path).lower()
        name = path.name.lower()
        if "\\cache\\" in p or "/cache/" in p:
            continue
        if "\\outputs\\revisao_valuation\\" in p or "/outputs/revisao_valuation/" in p:
            continue
        if "valuation" in name or "valuations" in p or "data_quality" in p or name.startswith("audit_"):
            valuation_files.append(path)
    return sorted(valuation_files, key=lambda p: (p.suffix.lower(), str(p).lower()))


def is_deep_excel(path: Path) -> bool:
    p = str(path).lower()
    return "\\outputs\\valuations\\" in p or "/outputs/valuations/" in p


def quick_inventory_xlsx(path: Path, findings: list[Finding]) -> list[dict[str, Any]]:
    """Fast inventory for legacy snapshots that are not part of the canonical output tree."""
    rows: list[dict[str, Any]] = []
    try:
        with zipfile.ZipFile(path) as zf:
            xml = zf.read("xl/workbook.xml")
        root = ET.fromstring(xml)
        ns = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        sheets = root.findall(".//main:sheets/main:sheet", ns)
        for sheet in sheets:
            rows.append({
                "arquivo": rel_file(path),
                "aba": sheet.attrib.get("name", "-"),
                "linhas": "",
                "colunas": "",
                "formulas": "",
                "populated": "",
                "years": "",
                "labels": "",
                "observacao": "inventario rapido; snapshot legado fora de outputs/valuations",
            })
    except Exception as exc:
        rows.append({
            "arquivo": rel_file(path),
            "aba": "-",
            "linhas": "",
            "colunas": "",
            "formulas": "",
            "populated": "",
            "years": "",
            "labels": "",
            "observacao": f"falha no inventario rapido: {exc}",
        })
    add_finding(
        findings, path, "workbook", "-", "baixa", "escopo",
        "Snapshot legado inventariado sem auditoria profunda",
        "O arquivo esta fora de `outputs/valuations` e foi tratado como snapshot historico para evitar que workbooks antigos bloqueiem a revisao do motor atual.",
        "O risco operacional principal esta nos arquivos canonicos atuais; snapshots antigos podem conter problemas ja superados.",
        "Regenerar ou arquivar snapshots antigos apos validar os canonicos.",
    )
    return rows


def year_columns(ws) -> dict[int, int]:
    columns: dict[int, int] = {}
    for row in range(1, min(ws.max_row, 12, MAX_SCAN_ROWS) + 1):
        for col in range(1, min(ws.max_column, MAX_SCAN_COLS) + 1):
            value = ws.cell(row=row, column=col).value
            if isinstance(value, int) and 1990 <= value <= CURRENT_YEAR + 20:
                columns[value] = col
            elif isinstance(value, str) and re.fullmatch(r"\d{4}", value.strip()):
                year = int(value.strip())
                if 1990 <= year <= CURRENT_YEAR + 20:
                    columns[year] = col
    return dict(sorted(columns.items()))


def label_rows(ws) -> dict[str, tuple[int, int, str]]:
    labels: dict[str, tuple[int, int, str]] = {}
    has_page_nav = norm(ws.cell(row=1, column=1).value) == "paginas"
    for row in range(1, min(ws.max_row, MAX_SCAN_ROWS) + 1):
        for col in range(1, min(ws.max_column, 8, MAX_SCAN_COLS) + 1):
            if has_page_nav and col <= 2:
                continue
            value = ws.cell(row=row, column=col).value
            if isinstance(value, str) and value.strip():
                labels.setdefault(norm(value), (row, col, value.strip()))
    return labels


def find_label(labels: dict[str, tuple[int, int, str]], patterns: list[str]) -> tuple[int, int, str] | None:
    pattern_norms = [norm(p) for p in patterns]
    for key, item in labels.items():
        for pat in pattern_norms:
            if pat and (pat == key or pat in key):
                return item
    return None


def row_values(ws_values, row: int, columns: dict[int, int]) -> dict[int, float | None]:
    out = {}
    for year, col in columns.items():
        out[year] = as_float(ws_values.cell(row=row, column=col).value)
    return out


def add_finding(findings: list[Finding], path: Path, sheet: str, cell: str, gravidade: str,
                categoria: str, problema: str, explicacao: str, impacto: str, sugestao: str) -> None:
    findings.append(Finding(rel_file(path), sheet, cell, gravidade, categoria, problema, explicacao, impacto, sugestao))


def check_formula_surface(path: Path, ws_formula, ws_values, findings: list[Finding]) -> dict[str, int]:
    populated = 0
    formulas = 0
    error_count = 0
    for row in ws_formula.iter_rows(max_row=min(ws_formula.max_row, MAX_SCAN_ROWS), max_col=min(ws_formula.max_column, MAX_SCAN_COLS)):
        for cell in row:
            value = cell.value
            if value is not None:
                populated += 1
            if isinstance(value, str) and value.startswith("="):
                formulas += 1
                if any(err in value.upper() for err in FORMULA_ERRORS):
                    error_count += 1
                    add_finding(
                        findings, path, ws_formula.title, cell.coordinate, "critica", "formula",
                        "Formula contem erro textual",
                        f"A formula contem marcador de erro: {value[:120]}",
                        "Pode invalidar diretamente o resultado da linha ou aba.",
                        "Revisar a referencia e regenerar a planilha apos corrigir o gerador.",
                    )
    for row in ws_values.iter_rows(max_row=min(ws_values.max_row, MAX_SCAN_ROWS), max_col=min(ws_values.max_column, MAX_SCAN_COLS)):
        for cell in row:
            value = cell.value
            if isinstance(value, str) and any(err in value.upper() for err in FORMULA_ERRORS):
                add_finding(
                    findings, path, ws_values.title, cell.coordinate, "critica", "formula",
                    "Celula calculada contem erro Excel",
                    f"O valor calculado retornou {value}.",
                    "Pode contaminar o valuation ou indicadores dependentes.",
                    "Abrir a formula, corrigir a referencia e recalcular o workbook.",
                )
    return {"populated": populated, "formulas": formulas, "error_count": error_count}


def check_sheet_structure(path: Path, ws_formula, ws_values, findings: list[Finding]) -> dict[str, Any]:
    labels = label_rows(ws_formula)
    years = year_columns(ws_formula)
    stats = check_formula_surface(path, ws_formula, ws_values, findings)
    stats.update({"years": ",".join(map(str, years)), "labels": len(labels)})

    sheet_norm = norm(ws_formula.title)
    critical_sheet = any(
        token in sheet_norm
        for token in ["dre", "dcf", "balanco", "wacc", "equity", "preco", "projec", "tir", "indice"]
    )
    if critical_sheet and stats["populated"] > 30 and stats["formulas"] == 0:
        add_finding(
            findings, path, ws_formula.title, "aba inteira", "alta", "vinculo",
            "Aba critica sem formulas",
            "A aba contem dados de valuation, mas nenhuma formula Excel. Os calculos foram escritos como valores fixos.",
            "Reduz auditabilidade, impede rastrear dependencias e torna a planilha vulneravel a atualizacoes parciais.",
            "Manter os resultados do Python, mas escrever tambem formulas nos blocos de DCF, WACC, multiplos e checagens.",
        )

    if years:
        ordered = list(years.keys())
        duplicates = [y for y, count in Counter(ordered).items() if count > 1]
        if duplicates:
            add_finding(
                findings, path, ws_formula.title, "cabecalho de anos", "alta", "datas",
                "Ano duplicado no cabecalho",
                f"Anos duplicados: {duplicates}.",
                "Valores historicos ou projetados podem ser gravados na coluna errada.",
                "Padronizar a linha de anos antes de escrever dados.",
            )
        gaps = [b for a, b in zip(ordered, ordered[1:]) if b != a + 1]
        if gaps and len(ordered) >= 4:
            add_finding(
                findings, path, ws_formula.title, "cabecalho de anos", "media", "datas",
                "Sequencia anual com lacunas ou colunas nao anuais",
                f"Cabecalho detectado: {ordered}.",
                "Pode indicar exclusao de ano historico, ponte trimestral residual ou desalinhamento entre historico e projecao.",
                "Usar apenas anos anuais na linha-base ou separar claramente colunas trimestrais/YTD/YTG.",
            )

    critical_patterns = [
        "receita", "ebitda", "ebit", "lucro liquido", "ativo total", "passivo total",
        "patrimonio liquido", "capital de giro", "capex", "depreciacao", "fcff", "fcfe",
        "divida liquida", "wacc", "ke custo", "valor do equity", "enterprise value",
        "preco justo", "valor da perpetuidade", "vp fcff", "vp fcfe",
    ]
    for pattern in critical_patterns:
        item = find_label(labels, [pattern])
        if not item or not years:
            continue
        row, _, original = item
        if row <= 2:
            continue
        vals = row_values(ws_values, row, years)
        non_blank = [v for v in vals.values() if v is not None]
        if not non_blank:
            add_finding(
                findings, path, ws_formula.title, f"{row}:{row}", "alta", "dados",
                f"Linha critica sem valores: {original}",
                "A linha foi identificada como item financeiro relevante, mas esta vazia nas colunas anuais.",
                "Pode deixar DRE/BP/DCF incompleto e distorcer margens, fluxo de caixa ou valuation.",
                "Conferir mapeamento da conta no escritor e garantir escrita nas colunas anuais corretas.",
            )
        elif all(abs(v) < 1e-12 for v in non_blank) and pattern not in {"divida liquida"}:
            add_finding(
                findings, path, ws_formula.title, f"{row}:{row}", "media", "dados",
                f"Linha critica zerada: {original}",
                "Todos os valores detectados na linha critica sao zero.",
                "Pode ser dado real, mas frequentemente indica falha de mapeamento ou escala.",
                "Validar contra o DataFrame normalizado e relatorio de completude CVM.",
            )
    return stats


def check_relations(path: Path, wb_values, wb_formula, findings: list[Finding]) -> None:
    for ws in wb_values.worksheets:
        wsf = wb_formula[ws.title]
        labels = label_rows(wsf)
        years = year_columns(wsf)
        if not years:
            continue

        ativo = find_label(labels, ["ativo total"])
        passivo = find_label(labels, ["passivo total"])
        if ativo and passivo:
            ativo_vals = row_values(ws, ativo[0], years)
            passivo_vals = row_values(ws, passivo[0], years)
            for year in years:
                a = ativo_vals.get(year)
                p = passivo_vals.get(year)
                if a is None or p is None:
                    continue
                diff = abs(a - p)
                base = max(abs(a), abs(p), 1.0)
                if diff > max(1.0, base * 0.01):
                    add_finding(
                        findings, path, ws.title, f"{rel_addr(ativo[0], years[year])}/{rel_addr(passivo[0], years[year])}",
                        "alta", "balanco",
                        "Ativo total diferente de passivo total",
                        f"{year}: ativo={a:,.2f}, passivo={p:,.2f}, diferenca={diff:,.2f}.",
                        "Indica BP desequilibrado ou escrita em coluna/linha errada.",
                        "Reconciliar ativo, passivo e patrimonio liquido no normalizador e escritor do BP.",
                    )

        yoy = find_label(labels, ["crescimento yoy", "crescimento receita"])
        if yoy:
            prev_row = yoy[0] - 1
            vals = row_values(ws, prev_row, years)
            yoy_vals = row_values(ws, yoy[0], years)
            ordered = sorted(years)
            for year in ordered[1:]:
                cur = vals.get(year)
                prev = vals.get(year - 1)
                stated = yoy_vals.get(year)
                if cur is None or prev in (None, 0) or stated is None:
                    continue
                expected = cur / prev - 1
                if abs(expected - stated) > 0.015:
                    add_finding(
                        findings, path, ws.title, rel_addr(yoy[0], years[year]), "media", "formula",
                        "Crescimento YoY divergente da linha base",
                        f"{year}: esperado {expected:.2%}, encontrado {stated:.2%}.",
                        "Pode distorcer leitura de crescimento historico/projetado e narrativa da tese.",
                        "Escrever formula YoY vinculada a linha imediatamente superior ou corrigir valor calculado.",
                    )

        fcfe = find_label(labels, ["fcfe fluxo de caixa ao acionista", "fcfe"])
        lucro = find_label(labels, ["lucro liquido controladores", "lucro liquido"])
        da = find_label(labels, ["depreciacao e amortizacao"])
        capex = find_label(labels, ["capex"])
        cap_reg = find_label(labels, ["capital regulatorio"])
        if fcfe and lucro and da and capex and cap_reg:
            series = {
                "fcfe": row_values(ws, fcfe[0], years),
                "lucro": row_values(ws, lucro[0], years),
                "da": row_values(ws, da[0], years),
                "capex": row_values(ws, capex[0], years),
                "cap_reg": row_values(ws, cap_reg[0], years),
            }
            for year in years:
                parts = [series[k].get(year) for k in ["lucro", "da", "capex", "cap_reg", "fcfe"]]
                if any(v is None for v in parts):
                    continue
                expected = series["lucro"][year] + series["da"][year] + series["capex"][year] + series["cap_reg"][year]
                actual = series["fcfe"][year]
                if abs(expected - actual) > max(2.0, abs(actual) * 0.02):
                    add_finding(
                        findings, path, ws.title, rel_addr(fcfe[0], years[year]), "alta", "dcf",
                        "FCFE nao reconcilia com componentes",
                        f"{year}: esperado {expected:,.2f}, encontrado {actual:,.2f}.",
                        "Afeta diretamente VP dos fluxos, valor terminal e preco-alvo.",
                        "Revisar sinal de CAPEX/capital regulatorio e formula do FCFE.",
                    )

        fcff = find_label(labels, ["fcff"])
        nopat = find_label(labels, ["nopat"])
        var_ncg = find_label(labels, ["variacao ncg", "variacao capital de giro"])
        if fcff and nopat and da and capex and var_ncg:
            series = {
                "fcff": row_values(ws, fcff[0], years),
                "nopat": row_values(ws, nopat[0], years),
                "da": row_values(ws, da[0], years),
                "capex": row_values(ws, capex[0], years),
                "var_ncg": row_values(ws, var_ncg[0], years),
            }
            for year in years:
                parts = [series[k].get(year) for k in ["nopat", "da", "capex", "var_ncg", "fcff"]]
                if any(v is None for v in parts):
                    continue
                expected = series["nopat"][year] + series["da"][year] + series["capex"][year] + series["var_ncg"][year]
                actual = series["fcff"][year]
                if abs(expected - actual) > max(2.0, abs(actual) * 0.02):
                    add_finding(
                        findings, path, ws.title, rel_addr(fcff[0], years[year]), "alta", "dcf",
                        "FCFF nao reconcilia com NOPAT, D&A, CAPEX e capital de giro",
                        f"{year}: esperado {expected:,.2f}, encontrado {actual:,.2f}.",
                        "Afeta diretamente EV, equity value e preco-alvo.",
                        "Revisar sinais de CAPEX/NCG e formula do FCFF.",
                    )

        wacc = find_label(labels, ["wacc"])
        ke = find_label(labels, ["ke custo do equity", "custo do capital proprio"])
        kd = find_label(labels, ["kd liquido", "kd bruto"])
        if wacc:
            wacc_vals = row_values(ws, wacc[0], years)
            for year, value in wacc_vals.items():
                if value is None:
                    continue
                if value > 1:
                    add_finding(
                        findings, path, ws.title, rel_addr(wacc[0], years[year]), "alta", "unidade",
                        "WACC parece estar em percentual inteiro",
                        f"{year}: WACC={value}. Valores percentuais deveriam estar como 0,10 para 10%.",
                        "Pode reduzir o valor terminal quase a zero se usado no DCF.",
                        "Normalizar taxas como decimais e aplicar apenas formatacao percentual.",
                    )
                elif value <= 0.02 or value >= 0.30:
                    add_finding(
                        findings, path, ws.title, rel_addr(wacc[0], years[year]), "media", "premissa",
                        "WACC fora de faixa usual",
                        f"{year}: WACC={value:.2%}.",
                        "Pode indicar premissa extrema ou problema de unidade/escala.",
                        "Validar taxa livre de risco, beta, ERP, custo da divida e estrutura de capital.",
                    )
        if ke and kd:
            # Only a plausibility check: Ke should usually not be below after-tax debt cost by a large margin.
            ke_vals = row_values(ws, ke[0], years)
            kd_vals = row_values(ws, kd[0], years)
            for year in years:
                kev = ke_vals.get(year)
                kdv = kd_vals.get(year)
                if kev is None or kdv is None or kev == 0 or kdv == 0:
                    continue
                if kev + 0.01 < kdv:
                    add_finding(
                        findings, path, ws.title, f"{rel_addr(ke[0], years[year])}/{rel_addr(kd[0], years[year])}",
                        "media", "premissa",
                        "Custo de equity abaixo do custo da divida",
                        f"{year}: Ke={kev:.2%}, Kd={kdv:.2%}.",
                        "Premissa pode inverter hierarquia de risco e distorcer WACC.",
                        "Validar beta, ERP, spread de divida e aliquota efetiva.",
                    )


def audit_excel(path: Path, findings: list[Finding]) -> list[dict[str, Any]]:
    inventory = []
    if path.suffix.lower() == ".xls":
        add_finding(
            findings, path, "-", "-", "baixa", "arquivo",
            "Arquivo .xls nao auditado por openpyxl",
            "O formato binario antigo nao e suportado pelo motor usado na auditoria.",
            "Fica fora da varredura de formulas e consistencia.",
            "Converter para .xlsx se ainda for relevante.",
        )
        return inventory
    try:
        wb_formula = openpyxl.load_workbook(path, data_only=False, read_only=False)
        wb_values = openpyxl.load_workbook(path, data_only=True, read_only=False)
    except PermissionError:
        add_finding(
            findings, path, "-", "-", "critica", "arquivo",
            "Arquivo bloqueado para leitura",
            "O Excel/OneDrive provavelmente esta com o arquivo aberto.",
            "A auditoria nao conseguiu validar formulas nem valores.",
            "Fechar o arquivo e rodar a auditoria novamente.",
        )
        return inventory
    except Exception as exc:
        add_finding(
            findings, path, "-", "-", "alta", "arquivo",
            "Falha ao abrir workbook",
            str(exc),
            "Arquivo ficou fora da auditoria tecnica.",
            "Verificar se o arquivo esta corrompido ou protegido.",
        )
        return inventory

    for ws_formula in wb_formula.worksheets:
        ws_values = wb_values[ws_formula.title]
        stats = check_sheet_structure(path, ws_formula, ws_values, findings)
        inventory.append({
            "arquivo": rel_file(path),
            "aba": ws_formula.title,
            "linhas": ws_formula.max_row,
            "colunas": ws_formula.max_column,
            **stats,
        })
    check_relations(path, wb_values, wb_formula, findings)
    wb_formula.close()
    wb_values.close()
    return inventory


def audit_csv(path: Path, findings: list[Finding]) -> dict[str, Any]:
    rows = 0
    cols = 0
    header: list[str] = []
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as fh:
            reader = csv.reader(fh)
            for i, row in enumerate(reader):
                if i == 0:
                    header = row
                    cols = len(row)
                rows += 1
    except UnicodeDecodeError:
        with path.open("r", encoding="latin1", newline="") as fh:
            reader = csv.reader(fh)
            for i, row in enumerate(reader):
                if i == 0:
                    header = row
                    cols = len(row)
                rows += 1
    except Exception as exc:
        add_finding(
            findings, path, "-", "-", "media", "arquivo",
            "Falha ao ler CSV",
            str(exc),
            "CSV ficou fora da auditoria.",
            "Verificar codificacao e delimitador.",
        )
    if rows <= 1:
        add_finding(
            findings, path, "-", "-", "media", "dados",
            "CSV vazio ou sem registros",
            "Arquivo CSV nao contem linhas de dados.",
            "Pode indicar auditoria/saida incompleta.",
            "Regenerar o CSV ou remover se for temporario.",
        )
    return {"arquivo": rel_file(path), "aba": "-", "linhas": rows, "colunas": cols, "formulas": 0, "populated": rows * cols, "years": "", "labels": len(header)}


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        return
    fieldnames: list[str] = []
    for row in rows:
        for key in row.keys():
            if key not in fieldnames:
                fieldnames.append(key)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def severity_sort(f: Finding) -> tuple[int, str, str]:
    order = {"critica": 0, "alta": 1, "media": 2, "baixa": 3}
    return (order.get(f.gravidade, 9), f.arquivo, f.aba)


def build_report(files: list[Path], inventory: list[dict[str, Any]], findings: list[Finding]) -> str:
    by_sev = Counter(f.gravidade for f in findings)
    by_cat = Counter(f.categoria for f in findings)
    excel_count = sum(1 for p in files if p.suffix.lower() in {".xlsx", ".xlsm", ".xls"})
    csv_count = sum(1 for p in files if p.suffix.lower() == ".csv")
    formula_sheets = sum(1 for row in inventory if int(row.get("formulas") or 0) > 0)
    total_sheets = sum(1 for row in inventory if row.get("aba") != "-")
    canonical = [p for p in files if "outputs\\valuations\\" in str(p).lower() or "outputs/valuations/" in str(p).lower()]

    lines = []
    lines.append("# RELATORIO_REVISAO_VALUATION")
    lines.append("")
    lines.append(f"Gerado em: {datetime.now():%Y-%m-%d %H:%M:%S}")
    lines.append("")
    lines.append("## Resumo executivo")
    lines.append("")
    lines.append(f"- Arquivos relacionados a valuation mapeados: **{len(files)}** ({excel_count} Excel, {csv_count} CSV).")
    lines.append(f"- Workbooks canonicos em `outputs/valuations`: **{len(canonical)}**.")
    lines.append(f"- Abas Excel auditadas: **{total_sheets}**.")
    lines.append(f"- Abas com pelo menos uma formula Excel: **{formula_sheets}**.")
    lines.append(f"- Achados: critica={by_sev.get('critica', 0)}, alta={by_sev.get('alta', 0)}, media={by_sev.get('media', 0)}, baixa={by_sev.get('baixa', 0)}.")
    lines.append("")
    if formula_sheets == 0 and total_sheets:
        lines.append("**Risco principal:** as planilhas de valuation auditadas estao sendo gravadas majoritariamente como valores fixos, sem formulas Excel rastreaveis nos blocos de DRE, BP, DCF, WACC e preco-alvo. Isso nao significa que o Python calculou tudo errado, mas reduz muito a auditabilidade e aumenta o risco de planilhas parciais ficarem incoerentes apos atualizacoes.")
        lines.append("")
    lines.append("## Metodologia")
    lines.append("")
    lines.append("- Inventario recursivo de `.xlsx`, `.xlsm`, `.xls` e `.csv` ligados a valuation.")
    lines.append("- Auditoria profunda aplicada aos workbooks canonicos em `outputs/valuations`; snapshots historicos em `outputs/Valuation_*.xlsx` foram inventariados por abas para evitar bloqueios e duplicidade de achados antigos.")
    lines.append("- Leitura de workbooks com `openpyxl` em modo formulas e valores calculados.")
    lines.append(f"- Varredura limitada a area operacional relevante: primeiras {MAX_SCAN_ROWS} linhas e {MAX_SCAN_COLS} colunas por aba, para evitar ranges fantasmas de formatacao.")
    lines.append("- Varredura de formulas quebradas, erros Excel, abas criticas sem formulas, anos duplicados/lacunosos e linhas financeiras vazias/zeradas.")
    lines.append("- Testes de consistencia para DRE/DCF, BP, FCFE, FCFF, YoY e WACC quando os rotulos estavam presentes.")
    lines.append("- Separacao entre erro objetivo, risco estrutural e premissa discutivel.")
    lines.append("")
    lines.append("## Contagem por categoria")
    lines.append("")
    for cat, count in by_cat.most_common():
        lines.append(f"- {cat}: {count}")
    lines.append("")
    lines.append("## Achados tecnicos")
    lines.append("")
    lines.append("| Gravidade | Arquivo | Aba | Celula/intervalo | Problema | Impacto | Correcao sugerida |")
    lines.append("|---|---|---|---|---|---|---|")
    for finding in sorted(findings, key=severity_sort)[:500]:
        lines.append(
            f"| {finding.gravidade} | `{finding.arquivo}` | `{finding.aba}` | `{finding.celula}` | "
            f"{finding.problema}: {finding.explicacao} | {finding.impacto} | {finding.sugestao} |"
        )
    if len(findings) > 500:
        lines.append("")
        lines.append(f"> Relatorio truncado nos 500 achados mais relevantes. O CSV completo esta em `outputs/revisao_valuation/revisao_valuation_findings.csv`.")
    lines.append("")
    lines.append("## Arquivos e abas auditados")
    lines.append("")
    lines.append("| Arquivo | Aba | Linhas | Colunas | Formulas | Anos detectados |")
    lines.append("|---|---:|---:|---:|---:|---|")
    for row in inventory[:800]:
        lines.append(
            f"| `{row.get('arquivo')}` | `{row.get('aba')}` | {row.get('linhas')} | {row.get('colunas')} | {row.get('formulas')} | {row.get('years', '')} |"
        )
    lines.append("")
    lines.append("## Correcoes automaticas")
    lines.append("")
    lines.append("Nenhuma planilha Excel foi alterada automaticamente nesta rodada. Os achados objetivos encontrados exigem mudanca no gerador das planilhas ou regeneracao dos arquivos; corrigir diretamente celulas isoladas quebraria o vinculo com o pipeline e poderia mascarar o erro de origem. Backups devem ser criados antes de qualquer regravacao em massa dos workbooks.")
    lines.append("")
    lines.append("Correcoes objetivas ja aplicadas no codigo do projeto durante a revisao:")
    lines.append("- `modules/valuation.py`: relacao PN/ON corrigida para `ON + PN * relacao`, TIR PN ajustada pela paridade e preco teto respeitando a relacao.")
    lines.append("- `main.py`: cotacao, par ON/PN e market cap do WACC passaram a usar relacao especifica da empresa.")
    lines.append("- `modules/05_valuation.py`: arquivo legado alinhado para nao manter a formula PN/ON invertida.")
    lines.append("")
    lines.append("## Proximos ajustes recomendados")
    lines.append("")
    lines.append("1. Transformar blocos criticos de DCF, WACC, multiplos e checagens em formulas Excel rastreaveis, mantendo os valores calculados pelo Python como controle.")
    lines.append("2. Adicionar uma aba `Checks` em cada valuation com reconciliacoes: BP fecha, FCFF/FCFE fecha, EV-to-equity fecha, preco por acao fecha, g < WACC/Ke.")
    lines.append("3. Separar claramente historico anual de colunas trimestrais/YTD/YTG, evitando misturar anos e pontes na mesma linha de cabecalho.")
    lines.append("4. Regenerar os workbooks canonicos em `outputs/valuations` depois de corrigir o gerador.")
    lines.append("5. Manter a auditoria de completude CVM/B3 como gate antes de gerar valuation final.")
    lines.append("")
    return "\n".join(lines)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    files = discover_files()
    findings: list[Finding] = []
    inventory: list[dict[str, Any]] = []

    for i, path in enumerate(files, start=1):
        print(f"[{i}/{len(files)}] {rel_file(path)}", flush=True)
        if path.suffix.lower() == ".csv":
            inventory.append(audit_csv(path, findings))
        elif path.suffix.lower() in {".xlsx", ".xlsm"} and not is_deep_excel(path):
            inventory.extend(quick_inventory_xlsx(path, findings))
        else:
            inventory.extend(audit_excel(path, findings))

    write_csv(OUTPUT_DIR / "revisao_valuation_inventory.csv", inventory)
    write_csv(OUTPUT_DIR / "revisao_valuation_findings.csv", [asdict(f) for f in findings])
    REPORT_PATH.write_text(build_report(files, inventory, findings), encoding="utf-8")
    print(f"Arquivos mapeados: {len(files)}")
    print(f"Abas/CSVs inventariados: {len(inventory)}")
    print(f"Achados: {len(findings)}")
    print(f"Relatorio: {REPORT_PATH}")


if __name__ == "__main__":
    main()
